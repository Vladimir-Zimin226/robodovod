"""Versioned, proposal-only workbook for the current process/role route."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from calculation.intake import PROCESS_DEFINITIONS
from calculation_contracts import PROCESS_DEMAND_UNITS, ProcessCode
from object_profiles import APPLICATION_OBJECT_TYPES, get_official_profile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = "project-workbook-v1"
# Additive proposal fields: older v1 books remain importable; missing new rows
# are treated as unknown, never as a silently confirmed policy.
OPTIONAL_ECONOMICS_FIELDS = {
    "robotizable_share", "residual_operations", "robots_per_control_post",
    "robots_per_day_technician", "rotation_factor", "technician_presence",
    "control_mode", "technician_purchase_mode", "technician_raas_mode",
    "technician_qualification_confirmed", "implementation_mode",
    "implementation_percent", "raas_mode", "raas_percent_monthly",
}
HEADERS = (
    "record_id",
    "parameter_code",
    "label",
    "value",
    "unit",
    "status",
    "source",
    "notes",
)
CSV_HEADERS = ("schema_version", "profile_code", "sheet", *HEADERS)
STATUSES = {"UNKNOWN", "DATA", "ASSUMPTION"}
# label, unit, type, minimum, maximum; every row is retained, even UNKNOWN.
FIELDS = {
    "Объект": {
        "name": ("Название объекта", "text", "text", None, None),
        "timezone": ("Местный часовой пояс IANA", "IANA", "timezone", None, None),
        "total_area_m2": ("Общая площадь объекта", "m2", "decimal", 0, 100000000),
        "active_area_m2": ("Активная площадь объекта", "m2", "decimal", 0, 100000000),
    },
    "Зоны": {
        "label": ("Название зоны", "text", "text", None, None),
        "constraints": (
            "Ограничения проходов, пола, груза и безопасности",
            "text",
            "text",
            None,
            None,
        ),
    },
    "Процессы": {
        "process_code": (
            "Код процесса (список в инструкции)",
            "code",
            "process",
            None,
            None,
        ),
        "zone_id": ("record_id зоны", "code", "id", None, None),
        "active": ("Описать активную операцию", "YES/NO", "bool", None, None),
        "demand": ("Поток за сутки", "dynamic", "decimal", 0, 10000000),
        "shifts": ("Смен в сутки", "shift", "integer", 1, 3),
        "hours": ("Часов в смене", "h", "decimal", 0, 24),
        "days": ("Рабочих дней в году", "day", "integer", 1, 366),
        "distance": ("Плечо в одну сторону", "m", "decimal", 0, 1000000),
        "batch": ("Единиц груза за рейс", "unit/trip", "integer", 1, 100000),
        "exchange_seconds": ("Погрузка и выгрузка за рейс", "s", "decimal", 0, 86400),
        "cleaning_frequency": (
            "Уборок указанной площади за сутки",
            "1/day",
            "decimal",
            0,
            100,
        ),
    },
    "Роли": {
        "role_code": ("Код роли (список в инструкции)", "code", "text", None, None),
        "process_ids": (
            "record_id процессов через запятую",
            "code",
            "text",
            None,
            None,
        ),
        "headcount": (
            "Численность роли без двойного учёта",
            "person",
            "integer",
            1,
            100000,
        ),
        "salary": (
            "Зарплата до удержаний на человека за месяц",
            "RUB/person/month",
            "decimal",
            0,
            100000000,
        ),
    },
    "Экономика": {
        "evaluation_date": ("Дата оценки", "YYYY-MM-DD", "date", None, None),
        "robotizable_share": ("Доля работы роли, доступная роботизации (0.65 = 65%)", "1", "decimal", 0, 1),
        "residual_operations": ("Ручные операции, которые остаются", "text", "text", None, None),
        "robots_per_control_post": ("Роботов на пост диспетчера", "robot/post", "decimal", 0.01, 100000),
        "robots_per_day_technician": ("Роботов на дневного техника", "robot/person", "decimal", 0.01, 100000),
        "rotation_factor": ("Коэффициент ротации для круглосуточного покрытия", "1", "decimal", 1, 100),
        "technician_presence": ("Присутствие техника DAY_WORKLOAD/EACH_SHIFT/VENDOR", "code", "tech_presence", None, None),
        "control_mode": ("Диспетчер: перевод/найм/подрядчик", "code", "staff_mode", None, None),
        "technician_purchase_mode": ("Техник при покупке: перевод/найм/подрядчик", "code", "staff_mode", None, None),
        "technician_raas_mode": ("Техник при RaaS: перевод/найм/подрядчик", "code", "staff_mode", None, None),
        "technician_qualification_confirmed": ("Квалификация переведённого техника подтверждена", "YES/NO", "bool", None, None),
        "implementation_mode": ("Внедрение: FIXED/PERCENT", "code", "price_mode", None, None),
        "implementation_percent": ("Внедрение, процент цены оборудования (10 = 10%)", "%", "decimal", 0, 100),
        "raas_mode": ("Тариф RaaS: FIXED/PERCENT", "code", "price_mode", None, None),
        "raas_percent_monthly": ("Месячный RaaS, процент цены робота (2 = 2%)", "%/month", "decimal", 0, 100),
        "horizon_years": ("Горизонт оценки", "year", "integer", 5, 15),
        "discount_rate": ("Ставка дисконтирования (0.15 = 15%)", "1", "decimal", 0, 1),
        "manual_units_per_shift": (
            "Ручная выработка",
            "unit/shift",
            "decimal",
            0,
            10000000,
        ),
        "control_headcount": ("Диспетчеры сейчас", "person", "integer", 0, 100000),
        "control_monthly_gross": (
            "Gross диспетчера в месяц",
            "RUB/person/month",
            "decimal",
            0,
            100000000,
        ),
        "technician_headcount": ("Техники сейчас", "person", "integer", 0, 100000),
        "technician_monthly_gross": (
            "Gross техника в месяц",
            "RUB/person/month",
            "decimal",
            0,
            100000000,
        ),
        "implementation_cost_total_gross": (
            "Внедрение, gross всего",
            "RUB",
            "decimal",
            0,
            1000000000000,
        ),
        "annual_service_per_robot_gross": (
            "Сервис робота, gross за год",
            "RUB/year",
            "decimal",
            0,
            1000000000000,
        ),
        "warranty_years": ("Гарантия", "year", "integer", 0, 15),
        "average_power_w": ("Средняя мощность робота", "W", "decimal", 0, 10000000),
        "shared_site_capital_gross": (
            "Общий CAPEX площадки, gross",
            "RUB",
            "decimal",
            0,
            1000000000000,
        ),
        "shared_annual_cost_gross": (
            "Общий OPEX площадки, gross за год",
            "RUB/year",
            "decimal",
            0,
            1000000000000,
        ),
        "raas_monthly_per_robot_gross": (
            "Тариф услуги, gross на робота в месяц",
            "RUB/robot/month",
            "decimal",
            0,
            1000000000000,
        ),
        "raas_contract_months": ("Длительность услуги", "month", "integer", 1, 180),
        "raas_infrastructure_owner": (
            "Кто владеет инфраструктурой CUSTOMER/VENDOR",
            "code",
            "owner",
            None,
            None,
        ),
        "start_seconds_from_midnight": (
            "Модельное начало смены от 00:00",
            "s",
            "integer",
            0,
            86399,
        ),
        "purchase_price": (
            "Предложенная цена — хранится; расчёт пока использует цену каталога",
            "RUB",
            "decimal",
            0,
            1000000000000,
        ),
        "currency": ("Валюта предложения", "code", "text", None, None),
        "vat_basis": ("НДС и состав предложения", "text", "text", None, None),
        **{
            code: (label, "YES/NO", "bool", None, None)
            for code, label in [
                (
                    "role_salaries_confirmed_as_monthly_gross",
                    "Предложение: зарплаты gross/месяц",
                ),
                (
                    "organizer_price_currency_rub_confirmed",
                    "Предложение: цена каталога в рублях",
                ),
                (
                    "initial_battery_in_robot_price_confirmed",
                    "Предложение: батарея включена в цену",
                ),
                (
                    "battery_replacements_in_service_confirmed",
                    "Предложение: замена батарей в сервисе",
                ),
                (
                    "raas_vendor_scope_confirmed",
                    "Предложение: состав RaaS включает оборудование/ПО/интеграцию/сервис",
                ),
            ]
        },
    },
}
KINDS = {"warehouse": "WAREHOUSE", "airport": "AIRPORT", "medical_facility": "CLINIC"}
DEFAULT_PROCESS = {
    "warehouse": "warehouse_receiving_shipping",
    "airport": "airport_baggage",
    "medical_facility": "clinic_medicines",
}
DEFAULT_ROLE = {
    "warehouse": "forklift_driver",
    "airport": "baggage_handler",
    "medical_facility": "sanitary",
}


def definitions(profile_code):
    profile = get_official_profile(profile_code)
    return {
        **FIELDS,
        "Паспорт": {
            p.parameter_code: (p.label, p.unit, p.data_type, p.min_value, p.max_value)
            for p in profile.parameters()
        },
    }


def interview_prompt(profile_code: str) -> str:
    profile = get_official_profile(profile_code)
    processes = {
        str(code): {
            "scope": str(d.scope),
            "kinds": [str(k) for k in d.allowed_kinds],
            "units": [str(PROCESS_DEMAND_UNITS[str(k)]) for k in d.allowed_kinds],
            "roles": [str(r) for r in d.roles],
        }
        for code, d in PROCESS_DEFINITIONS.items()
        if str(d.object_kind) == KINDS[profile.code]
    }
    return f"""Помоги заполнить приложенную книгу РОБОДОВОД {VERSION}, профиль {profile.code}.
Проведи интервью по 2–3 вопроса: объект и зоны/ограничения; процессы и потоки;
нагрузка и единицы; плечо В ОДНУ сторону, единиц за рейс, обмен; график;
роли без двойного учёта, месячная зарплата GROSS; экономические условия и источники.
Отдельно спроси, какая доля работы действительно роботизируется и что останется людям;
как покрываются диспетчер и техник (перевод/найм/подрядчик), сколько роботов
на пост/техника, нужна ли техника во все смены, коэффициент ротации.
Внедрение и RaaS задаются либо суммой FIXED, либо долей PERCENT от gross цены:
10 означает 10% внедрения, 2 означает 2% в месяц для RaaS; это не доли 0.1/0.02.
Не заполняй оба режима как действующие.
Коды покрытия диспетчера: TRANSFER, HIRE, EXISTING; техника при покупке:
TRANSFER, HIRE, EXISTING, CONTRACTOR; техника при RaaS также VENDOR.
Присутствие техника: DAY_WORKLOAD, EACH_SHIFT, VENDOR. Квалификацию перевода
не подтверждай за пользователя: оставь technician_qualification_confirmed UNKNOWN.
Затем уточни неизвестные паспортные параметры. Не рассчитывай парк, финансы или пригодность.
Не выдумывай число, цену, источник, подтверждение или обязательные данные.
UNKNOWN: пустое значение (не 0); DATA: явный ответ с источником; ASSUMPTION:
согласованное предложение с объяснением. Числа типового объекта — только ASSUMPTION.
Не ставь окончательное подтверждение пользователя; YES в коммерческих условиях — лишь предложение.
Не меняй листы, заголовки, коды, единицы, schema_version и profile_code. Не вставляй формулы/макросы.
Все строки полей сохраняй, неизвестные оставляй UNKNOWN. Сохранённая книга не подтверждает готовность закупки.
Для дополнительных зон/процессов/ролей копируй ВСЕ строки записи, меняя record_id,
связи zone_id/process_ids; одна общая роль хранится один раз со всеми связанными процессами.
Поток demand использует единицу выбранного процесса (например pallet/day, item/day, m2/day).
Для уборки demand — площадь одного прохода, cleaning_frequency — число проходов за сутки.
У неподдержанной операции сохраняй параметры, объясняй границу; не заменяй её другой формулой.
Верни заполненный XLSX в исходной структуре и перечень пропусков/допущений.
Если файлы недоступны, верни UTF-8 CSV: {",".join(CSV_HEADERS)}.
Одна строка соответствует одному полю; sheet — имя листа; неизвестное value пусто;
в каждой строке schema_version={VERSION}, profile_code={profile.code}. Это совместимый импорт.
Допустимые процессы, единицы и роли:\n{processes}
"""


def template_rows(profile_code: str, demo: bool = False) -> list[dict[str, Any]]:
    profile = get_official_profile(profile_code)
    code = DEFAULT_PROCESS[profile.code]
    demand_unit = str(
        PROCESS_DEMAND_UNITS[
            str(PROCESS_DEFINITIONS[ProcessCode(code)].allowed_kinds[0])
        ]
    )
    demo_values = {
        "Объект": {"name": "Типовой объект — пример", "timezone": "Asia/Sakhalin",
                   "total_area_m2": "20000", "active_area_m2": "10000"},
        "Зоны": {
            "label": "Основная зона",
            "constraints": "Проход, пол, безопасность и масса груза требуют обследования",
        },
        "Процессы": {
            "demand": "2000" if profile.code == "warehouse" else "220",
            "shifts": "2",
            "hours": "11",
            "days": "365",
            "distance": "120",
            "batch": "1",
            "exchange_seconds": "90",
        },
        "Роли": {"headcount": "25", "salary": "120000"},
        "Экономика": {
            "evaluation_date": "2026-09-26",
            "robotizable_share": "0.65",
            "residual_operations": "Контроль исключений, подготовка груза и нестандартные операции",
            "robots_per_control_post": "10",
            "robots_per_day_technician": "20",
            "rotation_factor": "2.2",
            "technician_presence": "DAY_WORKLOAD",
            "control_mode": "HIRE",
            "technician_purchase_mode": "HIRE",
            "technician_raas_mode": "VENDOR",
            "implementation_mode": "FIXED",
            "raas_mode": "FIXED",
            "horizon_years": "5",
            "discount_rate": "0.15",
            "manual_units_per_shift": "100",
            "control_headcount": "0",
            "control_monthly_gross": "100000",
            "technician_headcount": "0",
            "technician_monthly_gross": "120000",
            "implementation_cost_total_gross": "500000",
            "annual_service_per_robot_gross": "120000",
            "warranty_years": "1",
            "average_power_w": "1000",
            "shared_site_capital_gross": "0",
            "shared_annual_cost_gross": "0",
            "raas_monthly_per_robot_gross": "180000",
            "raas_contract_months": "60",
            "raas_infrastructure_owner": "VENDOR",
            "start_seconds_from_midnight": "32400",
            **{k: "YES" for k, v in FIELDS["Экономика"].items()
               if v[2] == "bool" and k != "technician_qualification_confirmed"},
        },
        "Паспорт": {
            p.parameter_code: str(p.default_value) for p in profile.parameters()
        },
    }
    structural = {
        "Процессы": {"process_code": code, "zone_id": "main", "active": "YES"},
        "Роли": {"role_code": DEFAULT_ROLE[profile.code], "process_ids": "operation"},
    }
    result = []
    for sheet, fields in definitions(profile.code).items():
        rid = (
            "operation"
            if sheet == "Процессы"
            else "staff"
            if sheet == "Роли"
            else "main"
        )
        for field, (label, unit, *_rest) in fields.items():
            value = structural.get(sheet, {}).get(field, "")
            status, source = (
                ("DATA", "Выбор структуры шаблона — проверить")
                if value
                else ("UNKNOWN", "")
            )
            if demo and field in demo_values.get(sheet, {}):
                value = demo_values[sheet][field]
                status = "ASSUMPTION"
                source = "Типовой пример организаторов / сценарное допущение; проверить на объекте"
            result.append(
                {
                    "record_id": rid,
                    "parameter_code": field,
                    "label": label,
                    "value": value,
                    "unit": demand_unit if unit == "dynamic" else unit,
                    "status": status,
                    "source": source,
                    "notes": "",
                    "sheet": sheet,
                }
            )
    return result


def build_workbook(
    profile_code: str, demo: bool = False, file_format: str = "xlsx"
) -> tuple[str, bytes]:
    profile = get_official_profile(profile_code)
    rows = template_rows(profile.code, demo)
    name = f"{profile.code}-{VERSION}-{'demo' if demo else 'blank'}.{file_format}"
    if file_format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.DictWriter(stream, fieldnames=CSV_HEADERS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(
            {**r, "schema_version": VERSION, "profile_code": profile.code} for r in rows
        )
        return name, stream.getvalue().encode("utf-8-sig")
    workbook = Workbook()
    workbook.active.title = "Версия"
    workbook.active.append(["schema_version", VERSION])
    workbook.active.append(["profile_code", profile.code])
    instructions = workbook.create_sheet("Инструкция")
    instructions.column_dimensions["A"].width = 110
    for line in interview_prompt(profile.code).splitlines():
        instructions.append([line])
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[row[0].row].height = max(
            30, min(240, 15 * (len(str(row[0].value)) // 90 + 1))
        )
    for title in definitions(profile.code):
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        for row in rows:
            if row["sheet"] == title:
                sheet.append([row[h] for h in HEADERS])
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234653")
        for column, width in zip("ABCDEFGH", [18, 48, 66, 30, 24, 18, 72, 40]):
            sheet.column_dimensions[column].width = width
        validation = DataValidation(type="list", formula1='"UNKNOWN,DATA,ASSUMPTION"')
        sheet.add_data_validation(validation)
        validation.add("F2:F3000")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return name, stream.getvalue()


def is_workbook(payload: bytes, file_format: str) -> bool:
    if file_format == "CSV":
        lines = payload.decode("utf-8-sig", errors="replace").splitlines()
        return bool(lines) and lines[0].split(",")[0].strip('"') == "schema_version"
    try:
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        from project_file_intake import IntakeError

        raise IntakeError("XLSX_INVALID") from exc
    try:
        return "Версия" in workbook.sheetnames
    finally:
        workbook.close()


def inspect_workbook(
    filename: str, payload: bytes, profile_code: str, file_format: str
):
    from project_file_intake import IntakeError, IntakeResult, _check_xlsx_archive

    profile = get_official_profile(profile_code)
    errors, warnings, rows = [], [], []

    def issue(code, message, sheet=None, row=None, field=None):
        if row is None and sheet and field:
            row = next(
                (
                    item["row"]
                    for item in rows
                    if item.get("sheet") == sheet
                    and item.get("parameter_code") == field
                ),
                None,
            )
        errors.append(
            {
                "code": code,
                "message": message,
                "sheet": sheet,
                "row": row,
                "parameter_code": field,
                "next_step": "Исправьте указанную строку по инструкции шаблона и повторите preview.",
            }
        )

    if file_format == "CSV":
        try:
            text = payload.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text, newline=""))
            if tuple(reader.fieldnames or ()) != CSV_HEADERS:
                raise IntakeError("CSV_HEADER_INVALID")
            for number, row in enumerate(reader, 2):
                if number > 3001:
                    raise IntakeError("CSV_ROW_LIMIT_EXCEEDED")
                if None in row or None in row.values():
                    issue("ROW_SHAPE", "Количество колонок не совпадает", row=number)
                if row.get("schema_version") != VERSION:
                    issue("VERSION_MISMATCH", "Неверная версия книги", row=number)
                if row.get("profile_code") != profile.code:
                    issue("PROFILE_MISMATCH", "Неверный тип объекта", row=number)
                rows.append({**row, "row": number})
        except UnicodeDecodeError as exc:
            raise IntakeError("CSV_ENCODING_INVALID") from exc
    else:
        _check_xlsx_archive(payload)
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
        try:
            expected = {"Версия", "Инструкция", *definitions(profile.code)}
            if set(workbook.sheetnames) != expected:
                issue(
                    "SHEETS_MISMATCH",
                    "Нужны все листы исходного шаблона; дополнительные листы не поддержаны",
                )
            meta = workbook["Версия"] if "Версия" in workbook.sheetnames else None
            if meta is not None and (
                meta.max_row != 2
                or meta.max_column != 2
                or meta.cell(1, 1).value != "schema_version"
                or meta.cell(2, 1).value != "profile_code"
            ):
                issue(
                    "METADATA_INVALID",
                    "Сохраните ровно две исходные строки метаданных",
                    "Версия",
                )
            if meta is None or meta.cell(1, 2).value != VERSION:
                issue("VERSION_MISMATCH", "Неверная версия книги", "Версия", 1)
            if meta is None or meta.cell(2, 2).value != profile.code:
                issue("PROFILE_MISMATCH", "Неверный тип объекта", "Версия", 2)
            for sheet in workbook:
                if sheet.max_row > 3000 or sheet.max_column > (
                    8 if sheet.title not in {"Версия", "Инструкция"} else 2
                ):
                    raise IntakeError("XLSX_DIMENSIONS_LIMIT_EXCEEDED")
                for row in sheet.iter_rows():
                    if any(c.data_type == "f" for c in row):
                        issue(
                            "FORMULA_FORBIDDEN",
                            "Формулы не исполняются и не принимаются",
                            sheet.title,
                            row[0].row,
                        )
                if sheet.title not in definitions(profile.code):
                    continue
                if (
                    tuple(c.value for c in next(sheet.iter_rows(min_row=1, max_row=1)))
                    != HEADERS
                ):
                    issue(
                        "HEADER_INVALID", "Не меняйте заголовки колонок", sheet.title, 1
                    )
                    continue
                for number, cells in enumerate(
                    sheet.iter_rows(min_row=2, values_only=True), 2
                ):
                    if all(value is None for value in cells):
                        continue
                    rows.append(
                        {
                            **dict(zip(HEADERS, cells)),
                            "sheet": sheet.title,
                            "row": number,
                        }
                    )
        finally:
            workbook.close()
    fields = definitions(profile.code)
    records = {}
    provenance = {}
    values = {}
    seen = set()
    digest = hashlib.sha256(payload).hexdigest()
    for row in rows:
        sheet, rid, field = (
            str(row.get(k) or "").strip()
            for k in ("sheet", "record_id", "parameter_code")
        )
        key = f"{sheet}.{rid}.{field}"
        if (
            sheet not in fields
            or field not in fields.get(sheet, {})
            or not re.fullmatch(r"[a-z][a-z0-9_-]{0,47}", rid)
        ):
            issue(
                "UNKNOWN_FIELD",
                "Неизвестный лист, код или record_id",
                sheet,
                row["row"],
                field,
            )
            continue
        if key in seen:
            issue(
                "DUPLICATE_FIELD",
                "Повтор кода в одной записи",
                sheet,
                row["row"],
                field,
            )
            continue
        seen.add(key)
        value = "" if row.get("value") is None else str(row["value"]).strip()
        status, source = (
            str(row.get("status") or "").strip(),
            str(row.get("source") or "").strip(),
        )
        if any(len(str(v or "")) > 4000 for v in row.values()):
            issue(
                "TEXT_LIMIT", "Текст превышает 4000 символов", sheet, row["row"], field
            )
        if any(
            str(v or "").lstrip().startswith(("=", "+", "@", "-"))
            and not re.fullmatch(r"-?[0-9]+(?:[.,][0-9]+)?", str(v or "").strip())
            for v in (row.get("value"), row.get("source"), row.get("notes"))
        ):
            issue(
                "FORMULA_FORBIDDEN",
                "Формулы и команды недопустимы",
                sheet,
                row["row"],
                field,
            )
        if (
            status not in STATUSES
            or (status == "UNKNOWN" and value)
            or (status != "UNKNOWN" and (not value or not source))
        ):
            issue(
                "STATUS_VALUE_MISMATCH",
                "UNKNOWN должен быть пустым; DATA/ASSUMPTION требуют значения и источника",
                sheet,
                row["row"],
                field,
            )
        _label, unit, kind, minimum, maximum = fields[sheet][field]
        if unit != "dynamic" and str(row.get("unit") or "") != unit:
            issue(
                "UNIT_MISMATCH", f"Единица должна быть {unit}", sheet, row["row"], field
            )
        if value:
            try:
                if kind in {"decimal", "integer", "number"}:
                    number = Decimal(value.replace(",", "."))
                    if not number.is_finite() or (
                        kind == "integer" and number != number.to_integral_value()
                    ):
                        raise ValueError()
                    if minimum is not None and number < Decimal(str(minimum)):
                        raise ValueError()
                    if maximum is not None and number > Decimal(str(maximum)):
                        raise ValueError()
                    value = format(number, "f")
                    if sheet == "Паспорт":
                        allowed = profile.parameters_by_code()[field].allowed_values
                        if allowed and number not in [Decimal(str(v)) for v in allowed]:
                            raise ValueError()
                elif (
                    kind == "bool"
                    and value not in {"YES", "NO"}
                    or kind == "id"
                    and not re.fullmatch(r"[a-z][a-z0-9_-]{0,47}", value)
                ):
                    raise ValueError()
                elif kind == "staff_mode" and value not in {"TRANSFER", "HIRE", "EXISTING", "CONTRACTOR", "VENDOR"}:
                    raise ValueError()
                elif kind == "staff_mode" and field == "control_mode" and value not in {"TRANSFER", "HIRE", "EXISTING"}:
                    raise ValueError()
                elif kind == "staff_mode" and field == "technician_purchase_mode" and value == "VENDOR":
                    raise ValueError()
                elif kind == "price_mode" and value not in {"FIXED", "PERCENT"}:
                    raise ValueError()
                elif kind == "tech_presence" and value not in {"DAY_WORKLOAD", "EACH_SHIFT", "VENDOR"}:
                    raise ValueError()
                elif kind == "timezone":
                    ZoneInfo(value)
                elif kind == "date":
                    date.fromisoformat(value)
                elif (
                    kind == "owner"
                    and value not in {"CUSTOMER", "VENDOR"}
                    or kind == "process"
                    and (
                        value not in PROCESS_DEFINITIONS
                        or str(PROCESS_DEFINITIONS[ProcessCode(value)].object_kind)
                        != KINDS[profile.code]
                    )
                ):
                    raise ValueError()
            except (InvalidOperation, ValueError, ZoneInfoNotFoundError):
                issue(
                    "VALUE_INVALID",
                    "Тип или диапазон значения неверны",
                    sheet,
                    row["row"],
                    field,
                )
        item = {
            "value": value or None,
            "status": status,
            "unit": row.get("unit"),
            "source": source,
            "notes": str(row.get("notes") or ""),
            "confirmed_by_user": False,
        }
        records.setdefault(sheet, {}).setdefault(rid, {})[field] = item
        values[key] = item
        provenance[key] = {
            "kind": "FILE",
            "status": status,
            "source_ref": source,
            "source": {
                "name": filename,
                "sha256": digest,
                "sheet": sheet,
                "row": row["row"],
                "cell": f"D{row['row']}",
            },
        }
    for sheet, definitions_ in fields.items():
        if sheet not in records:
            issue("RECORD_MISSING", "Обязательная запись листа отсутствует", sheet)
        for rid, record in records.get(sheet, {}).items():
            for field in definitions_.keys() - record.keys() - ({"total_area_m2", "active_area_m2"} if sheet == "Объект" else OPTIONAL_ECONOMICS_FIELDS if sheet == "Экономика" else set()):
                issue(
                    "FIELD_MISSING",
                    "Обязательная строка поля отсутствует; неизвестное храните как UNKNOWN",
                    sheet,
                    field=field,
                )
        if sheet in {"Объект", "Паспорт", "Экономика"} and set(
            records.get(sheet, {})
        ) != {"main"}:
            issue(
                "SINGLE_RECORD_REQUIRED", "Этот лист содержит одну запись main", sheet
            )
    required = []
    for rid, process in records.get("Процессы", {}).items():

        def read(field, record=process):
            return record.get(field, {}).get("value")

        code = read("process_code")
        zone = read("zone_id")
        if zone not in records.get("Зоны", {}):
            issue("ZONE_LINK", "Зона процесса отсутствует", "Процессы", field="zone_id")
        if code in PROCESS_DEFINITIONS:
            definition = PROCESS_DEFINITIONS[ProcessCode(code)]
            units = {
                str(PROCESS_DEMAND_UNITS[str(k)]) for k in definition.allowed_kinds
            }
            if process.get("demand", {}).get("unit") not in units:
                issue(
                    "DEMAND_UNIT",
                    f"Поток требует единицу из {sorted(units)}",
                    "Процессы",
                    field="demand",
                )
            supported = str(definition.scope) in {
                "TRANSPORT_CYCLE",
                "DELIVERY_CYCLE",
                "CLEANING_AREA",
            }
            if not supported:
                warnings.append(
                    {
                        "code": "UNSUPPORTED_FORMULA",
                        "message": f"{code}: параметры сохраняются, готового расчётного пути нет",
                    }
                )
            if read("active") == "YES":
                for field in [
                    "demand",
                    "shifts",
                    "hours",
                    "days",
                    *(
                        ["distance", "batch", "exchange_seconds"]
                        if str(definition.scope)
                        in {"TRANSPORT_CYCLE", "DELIVERY_CYCLE"}
                        else []
                    ),
                ]:
                    if not read(field):
                        required.append(f"Процессы.{rid}.{field}")
        if read("shifts") and read("hours"):
            try:
                if Decimal(read("shifts")) * Decimal(read("hours")) > 24:
                    issue(
                        "SCHEDULE_OVER_24H",
                        "Сумма смен превышает сутки",
                        "Процессы",
                        field="hours",
                    )
            except InvalidOperation:
                pass
    role_codes = set()
    if profile.code != "warehouse":
        warnings.append(
            {
                "code": "UNSUPPORTED_CATALOG_ROUTE",
                "message": "Паспорт и процессы сохраняются. Готовый путь Brain/каталога для этого объекта не заявлен; используйте форму и проверку применимости.",
            }
        )
    for rid, role in records.get("Роли", {}).items():
        code = role.get("role_code", {}).get("value")
        if code in role_codes:
            issue(
                "ROLE_DUPLICATE",
                "Общую роль храните одной записью",
                "Роли",
                field="role_code",
            )
        role_codes.add(code)
        for process_id in str(role.get("process_ids", {}).get("value") or "").split(
            ","
        ):
            process = records.get("Процессы", {}).get(process_id.strip())
            process_code = (process or {}).get("process_code", {}).get("value")
            if (
                process_code not in PROCESS_DEFINITIONS
                or code not in PROCESS_DEFINITIONS[ProcessCode(process_code)].roles
            ):
                issue(
                    "ROLE_LINK",
                    "Роль или связанный процесс несовместимы",
                    "Роли",
                    field="process_ids",
                )
    report = {
        "status": "INVALID" if errors else "VALID",
        "parameter_count": len(rows),
        "accepted_count": len(values),
        "error_count": len(errors),
        "warning_count": len(warnings),
        "errors": errors,
        "warnings": warnings,
        "required_inputs": required,
        "unknown_fields": [
            key for key, item in values.items() if item["status"] == "UNKNOWN"
        ],
        "sources": [
            {"field": key, **item}
            for key, item in values.items()
            if item["status"] != "UNKNOWN"
        ],
    }
    normalized = (
        {
            "schema_version": VERSION,
            "object_type": APPLICATION_OBJECT_TYPES[profile.code],
            "records": records,
            "file_source": {"name": filename, "sha256": digest},
            "confirmations_pending": True,
        }
        if not errors
        else None
    )
    return IntakeResult(
        file_format,
        profile.code,
        VERSION,
        filename,
        len(payload),
        digest,
        report,
        normalized,
        {"workbook": {"kind": "FILE", "source": {"name": filename, "sha256": digest}}},
        values,
        provenance,
    )


def _brain_projection(normalized, process_id, process):
    from brain_api import FIELD_UNITS

    fields = {}

    def propose(key, item):
        if item and item.get("value") is not None:
            fields[key] = {
                "value": item["value"],
                "unit": FIELD_UNITS[key],
                "raw_text": item.get("source", ""),
                "provenance": "expert_assumption"
                if item["status"] == "ASSUMPTION"
                else "user",
                "source_ref": normalized["file_source"],
                "confidence": 1,
                "confirmed_by_user": False,
            }

    fields["object_type"] = {
        "value": normalized["object_type"],
        "unit": FIELD_UNITS["object_type"],
        "provenance": "user",
        "confirmed_by_user": False,
    }
    object_row = normalized["records"]["Объект"]["main"]
    propose("total_area_m2", object_row.get("total_area_m2"))
    propose("active_area_m2", object_row.get("active_area_m2"))
    code = process["process_code"]["value"]
    fields["process_type"] = {
        "value": ("cleaning" if code == "warehouse_cleaning" else "transport")
        if code in {"warehouse_cleaning", "warehouse_receiving_shipping"}
        and process["active"]["value"] == "YES"
        else "unsupported",
        "unit": FIELD_UNITS["process_type"],
        "provenance": "user",
        "confirmed_by_user": False,
    }
    for workbook, brain in {
        "demand": "operations_per_day",
        "shifts": "shifts_count",
        "hours": "shift_hours",
        "days": "operating_days",
        "distance": "avg_distance_m",
        "batch": "units_per_trip",
        "exchange_seconds": "exchange_seconds",
        "cleaning_frequency": "cleaning_frequency_per_day",
    }.items():
        propose(brain, process.get(workbook))
    if "operations_per_day" in fields:
        fields["operations_per_day"]["unit"] = process["demand"]["unit"]
    zone = normalized["records"]["Зоны"][process["zone_id"]["value"]]
    propose("zone_label", zone.get("label"))
    propose("zone_constraints", zone.get("constraints"))
    role = next(
        (
            row
            for row in normalized["records"]["Роли"].values()
            if process_id
            in [
                item.strip()
                for item in str(row["process_ids"]["value"] or "").split(",")
            ]
        ),
        {},
    )
    propose("staff_headcount", role.get("headcount"))
    propose("monthly_gross_salary", role.get("salary"))
    propose(
        "manual_units_per_shift",
        normalized["records"]["Экономика"]["main"].get("manual_units_per_shift"),
    )
    return fields


def stage_brain_profile(project, normalized):
    """Append a proposal version inside the caller's transaction; never touch runs."""
    from brain_api import PROFILE_KEY, _current, _next, _state

    state = _state(project)
    process_id, process = next(iter(normalized["records"]["Процессы"].items()))
    code = process["process_code"]["value"]
    process_fields = {}
    for rid, row in normalized["records"]["Процессы"].items():
        # Brain selects by process code. Repeated codes across zones stay in the
        # workbook/form; they cannot be silently collapsed into one fleet.
        row_code = row["process_code"]["value"]
        if row_code in {"warehouse_receiving_shipping", "warehouse_cleaning"}:
            process_fields.setdefault(row_code, _brain_projection(normalized, rid, row))
    _next(
        state,
        _current(state, project.id),
        fields=_brain_projection(normalized, process_id, process),
        imported_workbook=normalized,
        process_fields=process_fields,
        selected_process=code
        if code in {"warehouse_receiving_shipping", "warehouse_cleaning"}
        else None,
        active_processes=[
            p["process_code"]["value"]
            for p in normalized["records"]["Процессы"].values()
            if p["active"]["value"] == "YES"
        ],
    )
    return {**project.profile, PROFILE_KEY: state, "file_intake_v2": normalized}
