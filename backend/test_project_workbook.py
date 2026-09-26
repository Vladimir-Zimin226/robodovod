from __future__ import annotations

import csv
import io
import zipfile
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from project_file_intake import MAX_FILE_BYTES, IntakeError, inspect_project_file
from project_workbook import CSV_HEADERS, build_workbook, interview_prompt


@pytest.mark.parametrize("profile", ["warehouse", "airport", "medical_facility"])
@pytest.mark.parametrize("demo", [False, True])
def test_workbook_csv_xlsx_roundtrip(profile, demo):
    xlsx = inspect_project_file(*build_workbook(profile, demo), profile)
    csv_result = inspect_project_file(*build_workbook(profile, demo, "csv"), profile)
    assert xlsx.valid and csv_result.valid
    assert xlsx.normalized_input["records"] == csv_result.normalized_input["records"]
    assert all(not item["confirmed_by_user"] for item in xlsx.parameter_values.values())
    passport = xlsx.normalized_input["records"]["Паспорт"]["main"]
    assert (
        len(passport)
        == {"warehouse": 42, "airport": 39, "medical_facility": 57}[profile]
    )
    assert all(
        item["status"] == ("ASSUMPTION" if demo else "UNKNOWN")
        for item in passport.values()
    )
    if not demo:
        assert (
            xlsx.normalized_input["records"]["Процессы"]["operation"]["demand"]["value"]
            is None
        )
        assert xlsx.report["required_inputs"]
    assert "2–3" in interview_prompt(profile) and "Не выдумывай" in interview_prompt(
        profile
    )


def csv_rows():
    return list(
        csv.DictReader(
            io.StringIO(build_workbook("warehouse", True, "csv")[1].decode("utf-8-sig"))
        )
    )


def encode(rows):
    stream = io.StringIO()
    writer = csv.DictWriter(stream, CSV_HEADERS)
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


@pytest.mark.parametrize(
    ("change", "code"),
    [
        ("version", "VERSION_MISMATCH"),
        ("duplicate", "DUPLICATE_FIELD"),
        ("unit", "DEMAND_UNIT"),
        ("range", "VALUE_INVALID"),
        ("extra", "UNKNOWN_FIELD"),
        ("missing", "FIELD_MISSING"),
        ("formula", "FORMULA_FORBIDDEN"),
        ("status", "STATUS_VALUE_MISMATCH"),
        ("schedule", "SCHEDULE_OVER_24H"),
    ],
)
def test_invalid_workbook_has_actionable_errors(change, code):
    rows = csv_rows()
    demand = next(
        row
        for row in rows
        if row["sheet"] == "Процессы" and row["parameter_code"] == "demand"
    )
    if change == "version":
        rows[0]["schema_version"] = "future-v9"
    elif change == "duplicate":
        rows.append(deepcopy(rows[0]))
    elif change == "unit":
        demand["unit"] = "kg/day"
    elif change == "range":
        demand["value"] = "10000001"
    elif change == "extra":
        rows[0]["parameter_code"] = "invented"
    elif change == "missing":
        rows.pop(0)
    elif change == "formula":
        demand["value"] = "=1+1"
    elif change == "status":
        demand["status"] = "UNKNOWN"
    else:
        next(
            row
            for row in rows
            if row["sheet"] == "Процессы" and row["parameter_code"] == "hours"
        )["value"] = "13"
    result = inspect_project_file("invalid.csv", encode(rows), "warehouse")
    assert not result.valid and result.normalized_input is None
    assert code in {issue["code"] for issue in result.report["errors"]}
    assert all(issue["next_step"] for issue in result.report["errors"])


def test_formula_macros_and_size_are_rejected():
    name, payload = build_workbook("warehouse", True)
    book = load_workbook(io.BytesIO(payload))
    book["Объект"]["D2"] = "=1+1"
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    result = inspect_project_file(name, stream.getvalue(), "warehouse")
    assert not result.valid and any(
        item["code"] == "FORMULA_FORBIDDEN" for item in result.report["errors"]
    )
    stream = io.BytesIO(payload)
    with zipfile.ZipFile(stream, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(IntakeError, match="MACROS"):
        inspect_project_file(name, stream.getvalue(), "warehouse")
    with pytest.raises(IntakeError, match="TOO_LARGE"):
        inspect_project_file("large.csv", b"x" * (MAX_FILE_BYTES + 1), "warehouse")
    with pytest.raises(IntakeError, match="ARCHIVE_LIMIT"):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("big", b"x" * 27000000)
        inspect_project_file(name, stream.getvalue(), "warehouse")


def test_unknown_pasport_not_required_but_unsupported_operation_retained():
    rows = csv_rows()
    for row in rows:
        if row["sheet"] == "Паспорт":
            row.update(value="", status="UNKNOWN", source="")
        if row["sheet"] == "Процессы" and row["parameter_code"] == "process_code":
            row["value"] = "warehouse_storage"
        if row["sheet"] == "Роли" and row["parameter_code"] == "role_code":
            row["value"] = "storekeeper"
    result = inspect_project_file("passport.csv", encode(rows), "warehouse")
    assert result.valid and any(
        item["code"] == "UNSUPPORTED_FORMULA" for item in result.report["warnings"]
    )
    assert (
        result.normalized_input["records"]["Процессы"]["operation"]["process_code"][
            "value"
        ]
        == "warehouse_storage"
    )


def test_xlsx_dimension_limit():
    name, payload = build_workbook("warehouse")
    book = load_workbook(io.BytesIO(payload))
    book["Зоны"].cell(3001, 1, "too-many")
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    with pytest.raises(IntakeError, match="DIMENSIONS_LIMIT"):
        inspect_project_file(name, stream.getvalue(), "warehouse")


@pytest.mark.parametrize(
    ("change", "code"),
    [
        ("extra", "SHEETS_MISMATCH"),
        ("missing", "SHEETS_MISMATCH"),
        ("metadata", "METADATA_INVALID"),
    ],
)
def test_workbook_sheet_and_metadata_structure(change, code):
    name, payload = build_workbook("warehouse")
    book = load_workbook(io.BytesIO(payload))
    if change == "extra":
        book.create_sheet("Extra")
    elif change == "missing":
        del book["Паспорт"]
    else:
        book["Версия"]["A1"] = "renamed"
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    result = inspect_project_file(name, stream.getvalue(), "warehouse")
    assert not result.valid and code in {
        item["code"] for item in result.report["errors"]
    }


def test_brain_preserves_processes_and_does_not_substitute_unsupported_formula():
    import uuid
    from types import SimpleNamespace

    from project_workbook import stage_brain_profile

    rows = csv_rows()
    additions = []
    for row in rows:
        if row["sheet"] not in {"Процессы", "Роли"}:
            continue
        new = deepcopy(row)
        new["record_id"] = "cleaning" if row["sheet"] == "Процессы" else "cleaners"
        if row["sheet"] == "Процессы":
            if row["parameter_code"] == "process_code":
                new["value"] = "warehouse_cleaning"
            if row["parameter_code"] == "demand":
                new.update(value="10000", unit="m2/day")
            if row["parameter_code"] == "cleaning_frequency":
                new.update(value="2", status="DATA", source="Synthetic interview")
        else:
            if row["parameter_code"] == "role_code":
                new["value"] = "cleaner"
            if row["parameter_code"] == "process_ids":
                new["value"] = "cleaning"
        additions.append(new)
    result = inspect_project_file("multi.csv", encode(rows + additions), "warehouse")
    assert result.valid, result.report
    project = SimpleNamespace(id=uuid.uuid4(), profile={})
    staged = stage_brain_profile(project, result.normalized_input)
    current = staged["brain_v1"]["versions"][-1]
    assert (
        current["process_fields"]["warehouse_cleaning"]["cleaning_frequency_per_day"][
            "value"
        ]
        == "2"
    )
    assert (
        current["process_fields"]["warehouse_cleaning"]["operations_per_day"]["unit"]
        == "m2/day"
    )
    assert current["imported_workbook"] == result.normalized_input
    for row in rows:
        if row["sheet"] == "Процессы" and row["parameter_code"] == "process_code":
            row["value"] = "warehouse_storage"
        if row["sheet"] == "Роли" and row["parameter_code"] == "role_code":
            row["value"] = "storekeeper"
    result = inspect_project_file("unsupported.csv", encode(rows), "warehouse")
    assert result.valid
    assert (
        stage_brain_profile(project, result.normalized_input)["brain_v1"]["versions"][
            -1
        ]["fields"]["process_type"]["value"]
        == "unsupported"
    )
