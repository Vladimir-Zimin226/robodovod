from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import replace

import jsonschema
from calculation.evidence_export import EvidenceRunSnapshotV1
from calculation.final_export import build_final_export
from calculation.scheduling import (
    SimulationReportV1,
    SimulationRequestV1,
    run_simulation,
)
from calculation.service import analyze_capacity
from calculation_contracts import parse_capacity_analysis_request, semantic_digest
from economics_final import execute_economics_v3, execute_economics_v4
from economics_orchestrator import EconomicsExecutionContextV1, execute_economics_v2
from openpyxl import load_workbook
from pypdf import PdfReader
from simulation_artifacts import StoredSimulationEvidence
from test_economics_orchestrator import _capacity_request, _inputs, _snapshot
from test_readable_report import _full_runs

from scripts.build_f5_contracts import ROOT, expected_outputs


def _run(price: str | None = None, source: str | None = None,
         demand: str = "1000", *, presentation_v4: bool = False):
    run, linked = _full_runs()
    linked = linked.model_copy(update={"run_id": "00000000-0000-4000-8000-000000000025"})
    snapshot = _snapshot()
    request_raw = _capacity_request().model_dump(mode="json")
    request_raw["project_id"] = run.project_id
    request_raw["process"]["demand"]["raw_value"] = demand
    request_raw["process"]["demand"]["normalized_value"] = demand
    request = parse_capacity_analysis_request(request_raw)
    capacity = analyze_capacity(request, snapshot, linked.run_id)
    context = EconomicsExecutionContextV1(
        run_id=run.run_id, project_id=run.project_id, tenant_id="tenant.demo",
        capacity_request=request, capacity_response=capacity.response,
        constraint_report=capacity.constraints.model_dump(mode="json"),
        executability=capacity.executability.model_dump(mode="json"),
    )
    inputs = _inputs()
    if price is not None:
        inputs["purchase_price_override_gross"] = price
        inputs["purchase_price_source"] = source
    execution = (execute_economics_v4 if presentation_v4 else execute_economics_v3)(inputs, snapshot, context)
    raw = run.model_dump(mode="json")
    raw["input_snapshot"] = {"capacity_run_id": linked.run_id, "economics": inputs}
    raw["result_snapshot"] = execution.result_snapshot
    raw["scenario_spec_snapshot"] = execution.scenario_spec_snapshot
    raw["versions"]["application"] = execution.application_version
    for field, snapshot_key in (("input", "input_snapshot"), ("result", "result_snapshot"),
                                ("scenario_spec", "scenario_spec_snapshot")):
        raw["checksums"][field] = semantic_digest(raw[snapshot_key]).removeprefix("sha256:")
    linked_raw = linked.model_dump(mode="json")
    linked_raw["input_snapshot"] = request.model_dump(mode="json")
    linked_raw["result_snapshot"] = capacity.response.model_dump(mode="json")
    linked_raw["checksums"]["input"] = semantic_digest(linked_raw["input_snapshot"]).removeprefix("sha256:")
    linked_raw["checksums"]["result"] = semantic_digest(linked_raw["result_snapshot"]).removeprefix("sha256:")
    return EvidenceRunSnapshotV1.model_validate(raw), EvidenceRunSnapshotV1.model_validate(linked_raw), execution


def test_u4_new_presentation_matches_saved_values_across_pdf_csv_xlsx_zip():
    for demand in ("220", "1000"):
        old, linked, _ = _run(demand=demand)
        new, _, execution = _run(demand=demand, presentation_v4=True)
        assert execution.application_version == "production-economics-orchestrator-v4"
        assert new.result_snapshot["comparison"] == old.result_snapshot["comparison"]
        old_archive = build_final_export(old, linked).archive
        package = build_final_export(new, linked)
        assert build_final_export(old, linked).archive == old_archive
        assert "Обзор" not in load_workbook(io.BytesIO(build_final_export(old, linked).files["Результат.xlsx"]), read_only=True).sheetnames
        workbook = load_workbook(io.BytesIO(package.files["Результат.xlsx"]), read_only=True, data_only=True)
        rows = list(workbook["Обзор"].values)
        csv_text = package.files["Сравнение.csv"].decode("utf-8-sig")
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(package.files["Отчёт_Рободовод.pdf"])).pages)
        pdf_pages = PdfReader(io.BytesIO(package.files["Отчёт_Рободовод.pdf"])).pages
        assert all((int(page.mediabox.width), int(page.mediabox.height)) == (595, 842) for page in pdf_pages)
        assert not re.search(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", pdf_text)
        assert rows[0] == ("Раздел", "Сценарий", "Год", "Показатель/статья", "Значение", "Единица", "Источник/база")
        assert ("Процесс и парк", "Все сценарии", None, "Плечо маршрута", "120", "m", "подтверждённый ввод расчёта парка") in rows
        fleet = str(linked.result_snapshot["capacity"]["value"]["selected_fleet"])
        assert any(row[0] == "Процесс и парк" and row[3] == "Выбранный парк" and row[4] == fleet for row in rows)
        assert "Эффективная производительность" in pdf_text and "Процесс и парк" in csv_text
        saved = new.result_snapshot["comparison"]
        purchase = next(item for item in saved["scenarios"] if item["scenario_id"] == "scenario.purchase.base")
        for value in (purchase["metrics"]["npv"]["value"], purchase["annual_cashflows"][0]["scenario"], purchase["capital_lines"][0]["amount"]):
            grouped = (f"{float(value):,.2f}" if "." in value else f"{int(value):,}").replace(",", " ")
            assert value in csv_text and grouped in pdf_text and any(value in row for row in rows[1:])
        assert "Денежные потоки" in csv_text and "Статьи CAPEX" in pdf_text
        with zipfile.ZipFile(io.BytesIO(package.archive)) as archive:
            assert archive.read("Результат.xlsx") == package.files["Результат.xlsx"]
            assert archive.read("Сравнение.csv") == package.files["Сравнение.csv"]


def test_f5_run_has_three_base_scenarios_and_per_scenario_sensitivity():
    run, _, execution = _run()
    result = run.result_snapshot
    comparison = result["comparison"]
    schema = json.loads((ROOT / "contracts/commercial-scenarios-bundle-v3.schema.json").read_text(encoding="utf-8"))
    jsonschema.Draft202012Validator(schema).validate(result)
    assert result["schema_version"] == "commercial-scenarios-bundle-v3"
    assert execution.application_version == "production-economics-orchestrator-v3"
    assert comparison["horizon_years"] >= 5
    assert len(comparison["scenarios"]) == 6
    assert comparison["baseline"]["metrics"]["roi"]["status"] == "N_A"
    assert comparison["baseline"]["metrics"]["tco"]["value"]
    purchase = next(item for item in comparison["scenarios"] if item["scenario_id"] == "scenario.purchase.base")
    raas = next(item for item in comparison["scenarios"] if item["scenario_id"] == "scenario.raas.base")
    assert purchase["metrics"]["roi"]["basis"] == raas["metrics"]["roi"]["basis"]
    assert purchase["metrics"]["tco"]["value"] != raas["metrics"]["tco"]["value"]
    assert len(comparison["sensitivity"]["by_scenario"]) == 7
    assert all(len(values) == 6 for values in comparison["sensitivity"]["by_scenario"].values())
    assert {item["parameter"] for item in comparison["sensitivity"]["by_scenario"]["scenario.raas.base"]} == {
        "RAAS_TARIFF", "OPERATION_VOLUME", "ROLE_SALARY"}


def test_f5_price_override_has_source_and_creates_new_saved_result():
    original, _, _ = _run()
    changed, _, _ = _run("3300000.00", "Пользовательский сценарий цены от 26.09.2026")
    assert original.result_snapshot != changed.result_snapshot
    assert changed.result_snapshot["comparison"]["inputs"]["price"]["value"] == "3300000.00"
    assert "Пользовательский" in changed.result_snapshot["comparison"]["inputs"]["price"]["source_note"]
    assert changed.result_snapshot["scenarios"][1]["procurement"]["procurement_ready"] is False


def test_user_price_with_source_can_replace_missing_catalog_price(monkeypatch):
    monkeypatch.setattr("economics_orchestrator.catalog_commercial_money", lambda _position: None)
    run, _, _ = _run("3300000.00", "Коммерческое предложение пользователя от 26.09.2026")
    assert run.result_snapshot["comparison"]["inputs"]["price"]["currency"] == "RUB"
    assert run.result_snapshot["comparison"]["inputs"]["price"]["source_note"].startswith("Коммерческое")


def test_220_pallets_at_120_m_and_typical_warehouse_export_saved_values():
    for demand in ("220", "1000"):
        run, linked, _ = _run(demand=demand)
        comparison = run.result_snapshot["comparison"]
        assert comparison["inputs"]["demand"]["value"] == demand
        assert linked.input_snapshot["process"]["route_distance"]["normalized_value"] == "120"
        assert len(comparison["scenarios"]) == 6
        package = build_final_export(run, linked)
        assert package.manifest.source_snapshot_digests["result"] == semantic_digest(run.result_snapshot)
        assert len(list(csv.DictReader(io.StringIO(package.files["Сравнение.csv"].decode("utf-8-sig")), delimiter=";"))) == 84
        assert package.files["Отчёт_Рободовод.pdf"].startswith(b"%PDF-")


def test_f5_archive_pdf_xlsx_csv_share_one_verified_snapshot():
    run, linked, _ = _run()
    first = build_final_export(run, linked)
    second = build_final_export(run, linked)
    assert first.archive == second.archive
    assert first.manifest.schema_version == "calculation-evidence-export-manifest-v4"
    schema = json.loads((ROOT / "contracts/calculation-evidence-export-manifest-v4.schema.json").read_text(encoding="utf-8"))
    jsonschema.Draft202012Validator(schema).validate(first.manifest.model_dump(mode="json"))
    with zipfile.ZipFile(io.BytesIO(first.archive)) as archive:
        assert archive.testzip() is None
        assert set(archive.namelist()) == {item.filename for item in first.manifest.artifacts} | {"manifest.json"}
        assert json.loads(archive.read("manifest.json")) == first.manifest.model_dump(mode="json")
        for artifact in first.manifest.artifacts:
            assert artifact.sha256 == "sha256:" + hashlib.sha256(archive.read(artifact.filename)).hexdigest()
        assert archive.read("Отчёт_Рободовод.pdf") == first.files["Отчёт_Рободовод.pdf"]
    rows = list(csv.DictReader(io.StringIO(first.files["Сравнение.csv"].decode("utf-8-sig")), delimiter=";"))
    assert len(rows) == 7 * 12
    assert {row["Горизонт, лет"] for row in rows} == {str(run.result_snapshot["comparison"]["horizon_years"])}
    assert {row["Расчёт"] for row in rows} == {run.run_id}
    book = load_workbook(io.BytesIO(first.files["Результат.xlsx"]), read_only=True, data_only=True)
    assert book.sheetnames == ["Итог", "Входы", "Оборудование", "Сценарии", "Денежные потоки", "Чувствительность", "Источники", "Симуляция", "Ограничения и версии"]
    assert book["Итог"].max_row == len(rows) + 1
    workbook_rows = list(book["Итог"].values)
    assert [dict(zip(workbook_rows[0], (cell or "" for cell in row))) for row in workbook_rows[1:]] == rows
    purchase_capex = next(row for row in rows if row["Сценарий"] == "Покупка · Базовый" and row["Показатель"] == "CAPEX")
    saved_capex = next(item for item in run.result_snapshot["comparison"]["scenarios"] if item["scenario_id"] == "scenario.purchase.base")["metrics"]["capex"]
    assert (purchase_capex["Значение"], purchase_capex["Единица"], purchase_capex["База/формула"], purchase_capex["Источник"]) == (
        saved_capex["value"], saved_capex["unit"], saved_capex["basis"], saved_capex["source_ref"])
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(first.files["Отчёт_Рободовод.pdf"])).pages)
    assert "Полное сравнение экономики F5" in text
    assert "ROI = эффект" in text
    assert "Сохранённая симуляция" in text
    assert saved_capex["value"] in text


def test_old_run_export_never_invents_missing_financial_metrics():
    run, linked = _full_runs()
    before = run.model_dump(mode="json")
    package = build_final_export(run, linked)
    old_archive = package.archive
    _run(demand="220")
    catalog = _snapshot()
    changed_catalog = replace(catalog, version=replace(catalog.version, code="future-catalog-disposable"))
    assert changed_catalog.version.code == "future-catalog-disposable"
    assert build_final_export(run, linked).archive == old_archive
    assert run.model_dump(mode="json") == before
    csv_text = package.files["Сравнение.csv"].decode("utf-8-sig")
    assert "NOT_SAVED" in csv_text
    assert "Полная сравнительная таблица" in csv_text
    assert package.manifest.simulation_report_digest is None


def test_selected_saved_c23_is_bound_to_pdf_svg_and_manifest():
    run, linked, _ = _run()
    request = SimulationRequestV1.model_validate({
        "schema_version": "simulation-request-v1", "request_id": "simulation.f5.sample",
        "tenant_id": run.result_snapshot["tenant_id"], "project_id": run.project_id,
        "scenario_spec": run.scenario_spec_snapshot, "mode": "DAILY",
        "peak_factor": None, "sla": None, "resources": [],
        "limits": {"max_jobs_per_day": 10000, "max_fleet": 100,
                   "max_runtime_seconds": 60, "progress_event_batch": 1000},
    })
    report = run_simulation(request)
    assert isinstance(report, SimulationReportV1)
    simulation = StoredSimulationEvidence(
        artifact_id=uuid.uuid4(), analysis_run_id=uuid.UUID(run.run_id),
        project_id=uuid.UUID(run.project_id), request=request, report=report,
        request_digest=semantic_digest(request), report_digest=semantic_digest(report),
        scenario_spec_digest=semantic_digest(request.scenario_spec),
    )
    package = build_final_export(run, linked, simulation)
    assert package.manifest.simulation_report_digest == simulation.report_digest
    assert package.manifest.scenario_spec_digest == simulation.scenario_spec_digest
    svg = ET.fromstring(package.files["Схема_2D.svg"])
    metadata = json.loads(svg.find("{http://www.w3.org/2000/svg}metadata").text)
    assert metadata["report_digest"] == simulation.report_digest
    assert metadata["run_id"] == run.run_id
    pdf = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(package.files["Отчёт_Рободовод.pdf"])).pages)
    assert simulation.report_digest in pdf
    assert str(report.queue.measurement_jobs) in pdf
    assert "Точка передачи A" in pdf and "Точка передачи B" in pdf
    assert "Условная схема 2D" in pdf
    foreign = StoredSimulationEvidence(**{**simulation.__dict__, "analysis_run_id": uuid.uuid4()})
    import pytest
    from calculation.evidence_export import EvidenceExportIntegrityError
    with pytest.raises(EvidenceExportIntegrityError):
        build_final_export(run, linked, foreign)


def test_generated_f5_contracts_are_current():
    assert all(path.read_bytes() == payload for path, payload in expected_outputs().items())


def test_v2_engine_still_accepts_original_fixture_without_f5_fields():
    snapshot = _snapshot()
    request = _capacity_request()
    capacity = analyze_capacity(request, snapshot, "run.capacity.demo")
    context = EconomicsExecutionContextV1(
        "run.economics.demo", request.project_id, "tenant.demo", request, capacity.response,
        capacity.constraints.model_dump(mode="json"), capacity.executability.model_dump(mode="json"),
    )
    result = execute_economics_v2(_inputs(), snapshot, context)
    assert result.result_snapshot["schema_version"] == "commercial-scenarios-bundle-v2"
