"""F5 deterministic final package.  This module only reads verified snapshots."""

from __future__ import annotations

import csv
import html
import io
import json
import re
import zipfile
from datetime import datetime
from decimal import Decimal
from typing import Any

from calculation_contracts import semantic_digest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader, PdfWriter
from simulation_artifacts import StoredSimulationEvidence

from calculation.evidence_export import (
    READABLE_REPORT_FILENAME,
    EvidenceExportIntegrityError,
    EvidenceExportManifestV4,
    EvidenceExportPackage,
    EvidenceRunSnapshotV1,
    _artifact,
    _verify_snapshots,
    build_evidence_export_v3,
)
from calculation.readable_report import _pdf

METRICS = (
    ("capex", "CAPEX"), ("opex_year_1", "Затраты/год 1"),
    ("opex_change_year_1", "Изменение затрат/год 1"), ("fot_year_1", "ФОТ/год 1"),
    ("effect_year_1", "Эффект/год 1"), ("effect_total", "Эффект за горизонт"),
    ("net_benefit", "Чистый эффект после вложений"), ("simple_payback", "Простая окупаемость"),
    ("roi", "ROI на CAPEX"), ("tco", "TCO за горизонт"),
    ("npv", "NPV"), ("discounted_payback", "Дисконтированная окупаемость"),
)
SCENARIO_LABELS = {"BASELINE": "Без роботов", "PURCHASE": "Покупка", "RAAS": "Услуга RaaS"}
UNCERTAINTY_LABELS = {"BASE": "Базовый", "PESSIMISTIC": "Пессимистичный", "OPTIMISTIC": "Оптимистичный"}
_FORMULA = re.compile(r"^[\s\x00-\x1f]*[=+\-@]")
_NUMBER = re.compile(r"^-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?$")
REPORT_PRESENTATION_VERSION = "result-presentation-v1"
INPUT_LABELS = {"demand": "Объём работ", "manual_productivity": "Ручная выработка",
                "price": "Цена робота", "raas_tariff": "Тариф услуги"}
SENSITIVITY_LABELS = {"EQUIPMENT_PRICE": "Цена робота", "RAAS_TARIFF": "Тариф услуги",
                      "OPERATION_VOLUME": "Объём работ", "ROLE_SALARY": "Зарплата роли",
                      "MANUAL_PRODUCTIVITY": "Ручная выработка"}


def _new_presentation(run: EvidenceRunSnapshotV1) -> bool:
    return (run.versions.get("application") == "production-economics-orchestrator-v4"
            and run.result_snapshot.get("versions", {}).get("report_presentation") == REPORT_PRESENTATION_VERSION)


def _presentation_rows(comparison: dict[str, Any] | None,
                       result: dict[str, Any] | None = None,
                       linked: EvidenceRunSnapshotV1 | None = None) -> list[list[str]]:
    """One read-only table for the new PDF, workbook and CSV views."""
    if comparison is None:
        return []
    rows: list[list[str]] = []
    if linked is not None:
        process = linked.input_snapshot.get("process") or {}
        capacity = (linked.result_snapshot.get("capacity") or {}).get("value") or {}
        for key, label in (("demand", "Объём работ"), ("route_distance", "Плечо маршрута")):
            field = process.get(key) or {}
            rows.append(["Процесс и парк", "Все сценарии", "", label,
                         str(field.get("normalized_value") if field.get("normalized_value") is not None else "нет данных"),
                         str(field.get("unit") or ""), "подтверждённый ввод расчёта парка"])
        for key, label, unit in (("recommended_fleet", "Рекомендованный парк", "роботов"),
                                 ("selected_fleet", "Выбранный парк", "роботов"),
                                 ("nominal_capacity", "Номинальная производительность", ""),
                                 ("effective_capacity", "Эффективная производительность", ""),
                                 ("coverage", "Покрытие объёма", "доля"),
                                 ("raw_load_ratio", "Фактическая загрузка", "доля")):
            raw = capacity.get(key)
            value = raw.get("value") if isinstance(raw, dict) else raw
            rows.append(["Процесс и парк", "Все сценарии", "", label,
                         str(value if value is not None else "нет данных"),
                         str(raw.get("unit") or unit) if isinstance(raw, dict) else unit,
                         "сохранённый расчёт парка"])
    for key, item in comparison.get("inputs", {}).items():
        rows.append(["Исходные данные", "Все сценарии", "", INPUT_LABELS.get(key, "Условие сценария"),
                     str(item.get("value") if item.get("value") is not None else "нет данных"),
                     str(item.get("unit") or item.get("currency") or ""),
                     str(item.get("source_note") or "подтверждённый ввод и сохранённый расчёт")])
    for scenario in _scenarios(comparison):
        name = _label(scenario)
        for key, label in METRICS:
            metric = _metric(scenario, key)
            rows.append(["Показатели", name, "", label,
                         str(metric.get("value") if metric.get("value") is not None else "нет данных"),
                         str(metric.get("unit") or ""), str(metric.get("basis") or "нет данных")])
        for flow in scenario.get("annual_cashflows", []):
            for field, label in (("baseline", "Без роботов"), ("scenario", "Сценарий"), ("effect", "Изменение")):
                rows.append(["Денежные потоки", name, str(flow.get("year") or ""), label,
                             str(flow.get(field) if flow.get(field) is not None else "нет данных"),
                             "RUB", "сохранённый денежный расчёт"])
        for line in scenario.get("capital_lines", []):
            rows.append(["Статьи CAPEX", name, "", str(line.get("label") or "Статья вложений").replace("_", " "),
                         str(line.get("amount") if line.get("amount") is not None else "нет данных"),
                         str(line.get("unit") or "RUB"), str(line.get("basis") or "нет данных")])
        detail = next((item for item in (result or {}).get("scenarios", [])
                       if item.get("acquisition") == scenario["acquisition"]
                       and item.get("uncertainty") == scenario["uncertainty"]), None)
        for line in (detail or {}).get("expenses", []):
            rows.append(["Расходы по статьям", name, "", str(line.get("label") or "Статья расходов"),
                         str(line.get("amount") if line.get("amount") is not None else "нет данных"),
                         "RUB", "сохранённая ведомость расходов"])
        for variant in comparison.get("sensitivity", {}).get("by_scenario", {}).get(scenario["scenario_id"], []):
            rows.append(["Чувствительность", name, "", f"{SENSITIVITY_LABELS.get(variant.get('parameter'), 'Условие')} {'−10%' if variant.get('direction') == 'LOWER' else '+10%'}",
                         str(variant.get("delta_npv") if variant.get("delta_npv") is not None else "нет данных"),
                         "RUB", str(variant.get("reason") or "повторный сохранённый расчёт")])
    for item in comparison.get("limitations", []):
        rows.append(["Ограничения", "Все сценарии", "", str(item), "", "", "сохранённый результат"])
    return rows


def _safe(value: Any) -> str:
    value = "" if value is None else str(value)
    return "'" + value if _FORMULA.match(value) and not _NUMBER.fullmatch(value) else value


def _display(value: str, unit: str) -> str:
    if _NUMBER.fullmatch(value):
        number = Decimal(value)
        grouped = f"{number:,.2f}" if "." in value else f"{number:,}"
        value = grouped.replace(",", " ")
    return f"{value} {unit.replace('RUB', '₽')}".strip()


def _comparison(run: EvidenceRunSnapshotV1) -> dict[str, Any] | None:
    result = run.result_snapshot
    if result.get("schema_version") != "commercial-scenarios-bundle-v3":
        return None
    projection = result.get("comparison")
    if not isinstance(projection, dict) or projection.get("schema_version") != "financial-comparison-v1":
        raise EvidenceExportIntegrityError("F5 run is missing its saved comparison")
    if projection.get("horizon_years", 0) < 5 or len(projection.get("scenarios", [])) != 6:
        raise EvidenceExportIntegrityError("F5 comparison is incomplete")
    if {item.get("scenario_id") for item in projection["scenarios"]} != {
        f"scenario.{acquisition}.{uncertainty}"
        for acquisition in ("purchase", "raas")
        for uncertainty in ("pessimistic", "base", "optimistic")
    }:
        raise EvidenceExportIntegrityError("F5 comparison scenario identities differ")
    expected_ids = {"scenario.baseline.base", *(item["scenario_id"] for item in projection["scenarios"])}
    if projection.get("baseline", {}).get("scenario_id") != "scenario.baseline.base":
        raise EvidenceExportIntegrityError("F5 baseline identity differs")
    sensitivity = projection.get("sensitivity", {})
    if sensitivity.get("schema_version") != "scenario-sensitivity-v1" or set(sensitivity.get("by_scenario", {})) != expected_ids:
        raise EvidenceExportIntegrityError("F5 scenario sensitivity is incomplete")
    for scenario_id, values in sensitivity["by_scenario"].items():
        if len(values) != 6 or len({(item.get("parameter"), item.get("direction")) for item in values}) != 6:
            raise EvidenceExportIntegrityError(f"F5 scenario sensitivity is incomplete: {scenario_id}")
    return projection


def _scenarios(comparison: dict[str, Any] | None) -> list[dict[str, Any]]:
    if comparison is None:
        return []
    return [comparison["baseline"], *comparison["scenarios"]]


def _label(row: dict[str, Any]) -> str:
    return f"{SCENARIO_LABELS.get(row['acquisition'], row['acquisition'])} · {UNCERTAINTY_LABELS.get(row['uncertainty'], row['uncertainty'])}"


def _metric(row: dict[str, Any], key: str) -> dict[str, Any]:
    value = row.get("metrics", {}).get(key)
    return value if isinstance(value, dict) else {
        "status": "NOT_SAVED", "value": None, "unit": "", "basis": "не сохранено в этой версии",
        "source_ref": "", "reason": "Для полной таблицы создайте новый расчёт",
    }


def _csv_rows(run: EvidenceRunSnapshotV1, comparison: dict[str, Any] | None) -> list[list[str]]:
    rows = [["Расчёт", "Сценарий", "Профиль", "Показатель", "Значение", "Единица", "Статус", "База/формула", "Источник", "Горизонт, лет", "Причина", "Источник цены"]]
    if comparison is None:
        rows.append([run.run_id, "Исторический/частичный", "", "Полная сравнительная таблица", "", "", "NOT_SAVED", "не сохранено в этой версии", "", "", "Создайте новый полный run", ""])
        return rows
    price = comparison.get("inputs", {}).get("price", {})
    price_source = price.get("source_note") or price.get("source_ref") or ""
    for scenario in _scenarios(comparison):
        for key, label in METRICS:
            metric = _metric(scenario, key)
            rows.append([run.run_id, _label(scenario), scenario["uncertainty"], label,
                         metric.get("value") or "", metric.get("unit") or "", metric.get("status") or "NOT_SAVED",
                         metric.get("basis") or "", metric.get("source_ref") or "",
                         str(comparison["horizon_years"]), metric.get("reason") or "", price_source])
    return rows


def comparison_csv(run: EvidenceRunSnapshotV1, comparison: dict[str, Any] | None,
                   linked: EvidenceRunSnapshotV1 | None = None) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerows([[_safe(value) for value in row] for row in _csv_rows(run, comparison)])
    if _new_presentation(run):
        writer.writerow([])
        writer.writerow(["Раздел", "Сценарий", "Год", "Показатель/статья", "Значение", "Единица", "Источник/база"])
        writer.writerows([[_safe(value) for value in row] for row in _presentation_rows(comparison, run.result_snapshot, linked)])
    return buffer.getvalue().encode("utf-8-sig")


def _sheet(book: Workbook, name: str, header: list[str], rows: list[list[Any]]) -> None:
    sheet = book.create_sheet(name)
    sheet.append(header)
    for row in rows:
        sheet.append([_safe(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163A37")
        cell.alignment = Alignment(wrap_text=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(55, max(14, *(len(str(cell.value or "")) + 2 for cell in column[:100])))


def comparison_xlsx(run: EvidenceRunSnapshotV1, linked: EvidenceRunSnapshotV1 | None,
                    comparison: dict[str, Any] | None,
                    simulation: StoredSimulationEvidence | None) -> bytes:
    book = Workbook()
    book.remove(book.active)
    book.properties.created = datetime(1980, 1, 1)
    book.properties.modified = datetime(1980, 1, 1)
    csv_rows = _csv_rows(run, comparison)
    _sheet(book, "Итог", csv_rows[0], csv_rows[1:])
    if _new_presentation(run):
        _sheet(book, "Обзор", ["Раздел", "Сценарий", "Год", "Показатель/статья", "Значение", "Единица", "Источник/база"], _presentation_rows(comparison, run.result_snapshot, linked))
    inputs = (comparison or {}).get("inputs", {})
    input_rows = [[name, item.get("value"), item.get("unit") or item.get("currency"),
                   item.get("tax_basis"), item.get("vat_rate"), item.get("source_ref"), item.get("source_note")]
                  for name, item in inputs.items()]
    if not input_rows:
        input_rows = [["Полные входы", "не сохранено в этой версии", "", "", "", "", ""]]
    _sheet(book, "Входы", ["Поле", "Значение", "Единица/валюта", "Налоговая база", "НДС", "Источник", "Примечание к источнику"], input_rows)
    equipment_rows = [[_label(s), line.get("label"), line.get("amount"), line.get("unit"),
                       line.get("basis"), ", ".join(line.get("source_refs", []))]
                      for s in _scenarios(comparison) for line in s.get("capital_lines", [])]
    _sheet(book, "Оборудование", ["Сценарий", "Статья CAPEX", "Сумма", "Единица", "База", "Источник"],
           equipment_rows or [["не сохранено в этой версии", "", "", "", "", ""]])
    _sheet(book, "Сценарии", ["Сценарий", "Профиль", "Горизонт", "Версия", "Статус"],
           [[_label(s), s["uncertainty"], comparison["horizon_years"], comparison["schema_version"], "СОХРАНЕНО"]
            for s in _scenarios(comparison)] if comparison else [["не сохранено в этой версии", "", "", "", "NOT_SAVED"]])
    flow_rows = [[_label(s), row.get("year"), row.get("baseline"), row.get("scenario"), row.get("effect"),
                  "RUB", row.get("source_ref")]
                 for s in _scenarios(comparison) for row in s.get("annual_cashflows", [])]
    _sheet(book, "Денежные потоки", ["Сценарий", "Год", "Без роботов", "Сценарий", "Изменение", "Единица", "Источник"],
           flow_rows or [["не сохранено в этой версии", "", "", "", "", "", ""]])
    sensitivity = (comparison or {}).get("sensitivity", {}).get("by_scenario", {})
    sens_rows = [[scenario, value.get("parameter"), value.get("direction"), value.get("base_value"),
                  value.get("variant_value"), value.get("unit"), value.get("baseline_npv"),
                  value.get("variant_npv"), value.get("delta_npv"), value.get("status"),
                  value.get("source_ref"), value.get("reason")]
                 for scenario, variants in sensitivity.items() for value in variants]
    _sheet(book, "Чувствительность", ["Сценарий", "Параметр", "Направление", "База", "Вариант ±10%", "Единица", "NPV базы", "NPV варианта", "Δ NPV", "Статус", "Источник", "Причина"],
           sens_rows or [["не сохранено в этой версии", "", "", "", "", "", "", "", "", "NOT_SAVED", "", ""]])
    sources = [["run", run.run_id], ["input", run.checksums.get("input")],
               ["result", run.checksums.get("result")], ["catalog", run.versions.get("catalog")],
               ["rules", run.versions.get("rules")], ["application", run.versions.get("application")]]
    if linked:
        sources.extend([["C11 run", linked.run_id], ["C11 result", linked.checksums.get("result")]])
    _sheet(book, "Источники", ["Тип", "Значение"], sources)
    simulation_rows = [["Статус", "Сохранённый C23 не выбран"], ["ScenarioSpec", ""], ["C23 digest", ""]]
    if simulation:
        report = simulation.report
        simulation_rows = [["Request", simulation.request.request_id], ["ScenarioSpec", simulation.scenario_spec_digest],
                           ["C23 digest", simulation.report_digest],
                           ["Задач измерения", report.queue.measurement_jobs],
                           ["Завершено", report.queue.completed_by_measurement_end],
                           ["SLA", report.sla.verdict], ["Загрузка", report.utilization.busy_fraction]]
    _sheet(book, "Симуляция", ["Показатель", "Значение"], simulation_rows)
    limitations = (comparison or {}).get("limitations", ["Полная сравнительная таблица не сохранена в этой версии."])
    _sheet(book, "Ограничения и версии", ["Тип", "Текст"],
           [["Ограничение", item] for item in limitations]
           + [["Версия", f"{key}: {value}"] for key, value in sorted(run.versions.items())])
    buffer = io.BytesIO()
    book.save(buffer)
    normalized = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9,
    ) as archive:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>1980-01-01T00:00:00Z\g<2>", payload,
                )
            archive.writestr(info, payload)
    return normalized.getvalue()


def visualization_svg(simulation: StoredSimulationEvidence) -> bytes:
    """Offline 2D schematic of the selected verified C23 report, never inferred KPI."""
    request = simulation.request
    report = simulation.report
    spec = request.scenario_spec
    meta = {"run_id": str(simulation.analysis_run_id), "request_id": request.request_id,
            "scenario_revision_id": spec.revision_id, "scenario_spec_digest": simulation.scenario_spec_digest,
            "request_digest": simulation.request_digest, "report_digest": simulation.report_digest,
            "simulation_time": "measurement_end", "geometry": "SCHEMATIC"}
    route = spec.routes[0].one_way_distance if spec.routes else None
    distance = f"{route.value} {route.unit}" if route else "не указан"
    utilization = (f"{Decimal(report.utilization.busy_fraction):.3f}"
                   if report.utilization.busy_fraction is not None else "не оценена")
    kpi = (f"Задач: {report.queue.measurement_jobs} · Завершено: {report.queue.completed_by_measurement_end}"
           f" · Маршрут: {distance} · SLA: {report.sla.verdict} · Загрузка: {utilization}")
    label = html.escape(str(spec.revision_id))
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="960" height="480" viewBox="0 0 960 480" role="img" aria-label="Схема 2D сохранённой симуляции">
<metadata>{html.escape(json.dumps(meta, ensure_ascii=False, sort_keys=True))}</metadata>
<rect width="960" height="480" fill="#f6faf8"/><text x="48" y="64" font-family="Arial" font-size="30" fill="#123c36">РОБОДОВОД · сохранённая симуляция 2D</text>
<text x="48" y="100" font-family="Arial" font-size="16" fill="#36554f">{label} · условная схема, не план объекта</text>
<rect x="68" y="160" width="215" height="145" rx="14" fill="#dceee5" stroke="#176950" stroke-width="3"/>
<text x="91" y="225" font-family="Arial" font-size="21" fill="#123c36">Точка передачи A</text>
<line x1="284" y1="232" x2="650" y2="232" stroke="#176950" stroke-width="13" stroke-linecap="round"/>
<polygon points="650,211 690,232 650,253" fill="#176950"/>
<rect x="690" y="160" width="215" height="145" rx="14" fill="#dceee5" stroke="#176950" stroke-width="3"/>
<text x="713" y="225" font-family="Arial" font-size="21" fill="#123c36">Точка передачи B</text>
<text x="48" y="380" font-family="Arial" font-size="18" fill="#123c36">{html.escape(kpi)}</text>
<text x="48" y="425" font-family="Arial" font-size="14" fill="#4d6660">Источник: проверенный C23 report. Геометрия схематическая; путь не является инженерным планом.</text>
</svg>'''
    return svg.encode("utf-8")


def final_pdf(run: EvidenceRunSnapshotV1, previous_pdf: bytes,
              comparison: dict[str, Any] | None,
              simulation: StoredSimulationEvidence | None,
              linked: EvidenceRunSnapshotV1 | None = None) -> bytes:
    if _new_presentation(run):
        lines: list[tuple[str, str]] = [
            ("РОБОДОВОД", "brand"), ("Экономика роботизации", "title"),
            (f"Сохранённый расчёт от {run.finished_at:%d.%m.%Y}", "subtitle"),
            ("Предварительная оценка. Цену, комплектацию и условия объекта следует подтвердить.", "cover"),
            ("", "page"),
        ]
        current = None
        for section, scenario, year, label, value, unit, source in _presentation_rows(comparison, run.result_snapshot, linked):
            heading = f"{section} · {scenario}"
            if heading != current:
                lines.append((heading, "section"))
                current = heading
            lines.append((f"{f'Год {year}: ' if year else ''}{label}: {_display(value, unit)}. Основание: {source}.", "body"))
        lines.append(("Проверка технических ограничений", "section"))
        lines.append(("Параметры поставщика, паспорт модели, рабочие условия площадки и состав поставки требуют отдельного подтверждения.", "body"))
        if simulation:
            lines.append(("Сохранённая симуляция", "section"))
            lines.append(("Симуляция связана с этим расчётом; схема условная и не является планом объекта.", "body"))
        lines.append(("Контрольные суммы и полная трассировка находятся в архиве ZIP.", "note"))
        return _pdf(lines)
    lines: list[tuple[str, str]] = [("Полное сравнение экономики F5", "section")]
    if comparison is None:
        lines.append(("Полная таблица baseline, покупки и услуги не сохранена в этой версии. Для неё нужен новый run.", "body"))
    else:
        lines.extend([("Значения взяты из сохранённого run. Прогноз предварительный; объект и условия поставки требуют обследования.", "note"),
                      (f"Общий горизонт: {comparison['horizon_years']} лет. Валюта: RUB.", "body")])
        price = comparison.get("inputs", {}).get("price", {})
        lines.append((f"Цена робота: {price.get('value') or 'не сохранено'} {price.get('currency') or 'RUB'}; база {price.get('tax_basis') or 'не указана'}; НДС {price.get('vat_rate') or 'не указан'}; источник {price.get('source_note') or price.get('source_ref') or 'не указан'}.", "body"))
        rows = [comparison["baseline"], *[item for item in comparison["scenarios"] if item["uncertainty"] == "BASE"]]
        for key, label in METRICS:
            values = []
            for scenario in rows:
                metric = _metric(scenario, key)
                value = metric.get("value") if metric.get("status") == "COMPLETE" else metric.get("status")
                values.append(f"{SCENARIO_LABELS[scenario['acquisition']]}: {value or 'не сохранено'} {metric.get('unit') or ''}")
            lines.append((f"{label}: {' | '.join(values)}", "body"))
        lines.extend([("ROI = эффект за горизонт / инвестиции в проект; при нулевом CAPEX — N/A. TCO = CAPEX + затраты за горизонт. NPV — отдельная дисконтированная метрика.", "note"),
                      ("Варианты неопределённости", "section")])
        for scenario in comparison["scenarios"]:
            lines.append((f"{_label(scenario)}: NPV {_metric(scenario, 'npv').get('value')} RUB; TCO {_metric(scenario, 'tco').get('value')} RUB; ROI {_metric(scenario, 'roi').get('value') or 'N/A'} %.", "body"))
        lines.append(("Чувствительность ±10% по выбранному сценарию", "section"))
        for scenario in rows:
            variants = comparison["sensitivity"]["by_scenario"].get(scenario["scenario_id"], [])
            lines.append((_label(scenario), "metric"))
            for variant in variants:
                lines.append((f"{variant['parameter']} {variant['direction']}: база {variant['base_value']} {variant['unit']}; Δ NPV {variant['delta_npv'] or variant['status']} RUB. {variant['reason'] or ''}", "body"))
        lines.append(("Источники и ограничения", "section"))
        for item in comparison["limitations"]:
            lines.append((item, "body"))
    lines.extend([("", "page"), ("Сохранённая симуляция и схема 2D", "section")])
    if simulation:
        report = simulation.report
        utilization = (f"{Decimal(report.utilization.busy_fraction):.3f}"
                       if report.utilization.busy_fraction is not None else "не оценена")
        lines.extend([(f"ScenarioSpec: {simulation.scenario_spec_digest}; C23: {simulation.report_digest}", "body"),
                      (f"Измерение: {report.queue.measurement_jobs} задач; завершено {report.queue.completed_by_measurement_end}; SLA: {report.sla.verdict}; загрузка: {utilization}.", "metric"),
                      ("Условная схема 2D; не является инженерным планом объекта. Отдельный SVG открывается без сервиса.", "diagram")])
    else:
        lines.append(("Сохранённый C23 для этого run не выбран или отсутствует. 2D и KPI не подменены расчётными значениями.", "body"))
    appendix = _pdf(lines)
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(previous_pdf)))
    writer.append(PdfReader(io.BytesIO(appendix)))
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_final_export(run: EvidenceRunSnapshotV1,
                       linked: EvidenceRunSnapshotV1 | None = None,
                       simulation: StoredSimulationEvidence | None = None) -> EvidenceExportPackage:
    """Verify every binding before serializing a single byte of the package."""
    if linked is not None and (linked.project_id != run.project_id
            or str(run.input_snapshot.get("capacity_run_id")) != linked.run_id):
        raise EvidenceExportIntegrityError("linked capacity run does not match saved economics run")
    _verify_snapshots(run)
    if linked is not None:
        _verify_snapshots(linked)
    if simulation is not None:
        if (str(simulation.analysis_run_id) != run.run_id
                or str(simulation.project_id) != run.project_id
                or run.scenario_spec_snapshot != simulation.request.scenario_spec.model_dump(mode="json")
                or f"sha256:{run.checksums.get('scenario_spec')}" != simulation.scenario_spec_digest):
            raise EvidenceExportIntegrityError("C23 does not match this run/ScenarioSpec")
    previous = build_evidence_export_v3(run, linked)
    comparison = _comparison(run)
    files = dict(previous.files)
    files["Сравнение.csv"] = comparison_csv(run, comparison, linked)
    files["Результат.xlsx"] = comparison_xlsx(run, linked, comparison, simulation)
    files[READABLE_REPORT_FILENAME] = final_pdf(run, files[READABLE_REPORT_FILENAME], comparison, simulation, linked)
    files["Связь_сценария.json"] = (json.dumps({
        "run_id": run.run_id, "result_digest": previous.manifest.source_snapshot_digests["result"],
        "scenario_spec_digest": simulation.scenario_spec_digest if simulation else (
            f"sha256:{run.checksums['scenario_spec']}" if run.checksums.get("scenario_spec") else None),
        "simulation_request_id": simulation.request.request_id if simulation else None,
        "simulation_request_digest": simulation.request_digest if simulation else None,
        "simulation_report_digest": simulation.report_digest if simulation else None,
        "visualization": "Схема_2D.svg" if simulation else None,
    }, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode("utf-8")
    if simulation:
        files["Схема_2D.svg"] = visualization_svg(simulation)
    files["НАЧНИТЕ_ЗДЕСЬ.md"] += ("\n## Полная таблица F5\n\nОткройте Результат.xlsx или Сравнение.csv. "
        "Денежные значения взяты из сохранённого run; строка NOT_SAVED не является нулём. "
        "Схема_2D.svg и Связь_сценария.json связаны с выбранным сохранённым C23.\n").encode()
    if _new_presentation(run):
        files["НАЧНИТЕ_ЗДЕСЬ.md"] += ("Новый отчёт содержит единый обзор: PDF «Таблицы сохранённого результата», лист «Обзор» XLSX и подробная часть CSV.\n").encode()
    def media(name: str) -> str:
        return ("application/pdf" if name.endswith(".pdf") else
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if name.endswith(".xlsx") else
                "text/csv; charset=utf-8" if name.endswith(".csv") else
                "image/svg+xml" if name.endswith(".svg") else
                "application/json" if name.endswith(".json") else "text/markdown; charset=utf-8")
    artifacts = [_artifact(name, media(name), payload) for name, payload in sorted(files.items())]
    body = previous.manifest.model_dump(mode="python")
    body["sections"] = previous.manifest.sections
    body.update(schema_version="calculation-evidence-export-manifest-v4",
                export_policy_version="calculation-evidence-export-policy-v4",
                generator_version="snapshot-evidence-export-v4",
                comparison_filename="Сравнение.csv", workbook_filename="Результат.xlsx",
                visualization_filename="Схема_2D.svg" if simulation else None,
                simulation_request_id=simulation.request.request_id if simulation else None,
                simulation_report_digest=simulation.report_digest if simulation else None,
                scenario_spec_digest=simulation.scenario_spec_digest if simulation else None,
                artifacts=artifacts, bundle_content_digest=semantic_digest([item.model_dump(mode="json") for item in artifacts]),
                limitations=[*previous.manifest.limitations, "financial-comparison-read-from-saved-run", "schematic-2d-not-engineering-plan"],
                manifest_digest="sha256:" + "0" * 64)
    body["manifest_digest"] = semantic_digest(EvidenceExportManifestV4.model_construct(**body).model_dump(mode="json"))
    manifest = EvidenceExportManifestV4.model_validate(body)
    manifest_payload = (json.dumps(manifest.model_dump(mode="json"), ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name, payload in sorted({**files, "manifest.json": manifest_payload}.items()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            archive.writestr(info, payload)
    return EvidenceExportPackage(manifest=manifest, archive=buffer.getvalue(), files=files)


__all__ = ["build_final_export", "comparison_csv", "comparison_xlsx", "final_pdf", "visualization_svg"]
