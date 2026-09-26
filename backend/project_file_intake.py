"""Strict XLSX/CSV adapters for project input preview and application."""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from object_profiles import (
    ObjectProfile,
    ObjectProfileError,
    ProfileParameter,
    build_profile_projection,
    get_official_profile,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_XLSX_EXPANDED_BYTES = 25 * 1024 * 1024
CSV_HEADERS = ("profile_code", "parameter_code", "value", "unit")
MEDIA_TYPES = {
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "CSV": "text/csv",
}


class IntakeError(ValueError):
    """A stable, non-sensitive file intake failure."""


@dataclass(frozen=True)
class IntakeResult:
    file_format: str
    profile_code: str
    profile_version: str
    original_name: str
    byte_size: int
    sha256: str
    report: dict[str, Any]
    normalized_input: dict[str, Any] | None
    provenance: dict[str, Any] | None
    parameter_values: dict[str, Any]
    parameter_provenance: dict[str, Any]

    @property
    def valid(self) -> bool:
        return not self.report["errors"]

    def public_dict(self) -> dict[str, Any]:
        return {
            "valid": self.valid,
            "file": {
                "name": self.original_name,
                "format": self.file_format,
                "media_type": MEDIA_TYPES[self.file_format],
                "byte_size": self.byte_size,
                "sha256": self.sha256,
            },
            "profile_code": self.profile_code,
            "profile_version": self.profile_version,
            "report": self.report,
            "normalized_input": self.normalized_input,
            "provenance": self.provenance,
            "parameter_values": self.parameter_values,
            "parameter_provenance": self.parameter_provenance,
        }


def detect_format(filename: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        if not payload.startswith(b"PK"):
            raise IntakeError("XLSX_SIGNATURE_INVALID")
        return "XLSX"
    if suffix == ".csv":
        return "CSV"
    raise IntakeError("FILE_TYPE_UNSUPPORTED")


def _issue(
    code: str,
    message: str,
    *,
    parameter_code: str | None = None,
    row: int | None = None,
) -> dict[str, Any]:
    value: dict[str, Any] = {"code": code, "message": message}
    if parameter_code is not None:
        value["parameter_code"] = parameter_code
    if row is not None:
        value["row"] = row
    return value


def _parse_number(raw: Any) -> int | float:
    if isinstance(raw, bool):
        raise ValueError
    if isinstance(raw, (int, float)):
        value = float(raw)
    elif isinstance(raw, str):
        normalized = raw.strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
        if not normalized:
            raise ValueError
        value = float(normalized)
    else:
        raise ValueError
    if not math.isfinite(value):
        raise ValueError
    return int(value) if value.is_integer() else value


def _validate_value(
    parameter: ProfileParameter,
    raw: Any,
    errors: list[dict[str, Any]],
    *,
    row: int,
) -> Any | None:
    code = parameter.parameter_code
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        errors.append(_issue("VALUE_MISSING", "Значение не указано", parameter_code=code, row=row))
        return None
    if parameter.data_type == "number":
        try:
            value = _parse_number(raw)
        except (TypeError, ValueError, OverflowError):
            errors.append(_issue("TYPE_MISMATCH", "Ожидается число", parameter_code=code, row=row))
            return None
        if parameter.min_value is not None and float(value) < float(parameter.min_value):
            errors.append(_issue("BELOW_MIN", f"Значение меньше {parameter.min_value}", parameter_code=code, row=row))
        if parameter.max_value is not None and float(value) > float(parameter.max_value):
            errors.append(_issue("ABOVE_MAX", f"Значение больше {parameter.max_value}", parameter_code=code, row=row))
    else:
        value = str(raw).strip()
    if (
        parameter.allowed_values
        and value != parameter.default_value
        and value not in parameter.allowed_values
    ):
        errors.append(_issue("VALUE_NOT_ALLOWED", "Значение отсутствует в допустимом списке", parameter_code=code, row=row))
    return value


def _parse_csv(
    payload: bytes, profile: ObjectProfile
) -> tuple[
    dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]
]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IntakeError("CSV_ENCODING_INVALID") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text, newline=""), dialect=dialect)
    if tuple(reader.fieldnames or ()) != CSV_HEADERS:
        raise IntakeError("CSV_HEADER_INVALID")
    parameters = profile.parameters_by_code()
    values: dict[str, Any] = {}
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    locations: dict[str, Any] = {}
    row_count = 0
    for row_number, row in enumerate(reader, start=2):
        row_count += 1
        if row_count > 500:
            raise IntakeError("CSV_ROW_LIMIT_EXCEEDED")
        row_profile = (row.get("profile_code") or "").strip()
        code = (row.get("parameter_code") or "").strip()
        if row_profile != profile.code:
            errors.append(_issue("PROFILE_MISMATCH", f"Ожидается профиль {profile.code}", parameter_code=code or None, row=row_number))
            continue
        parameter = parameters.get(code)
        if parameter is None:
            errors.append(_issue("UNKNOWN_PARAMETER", "Неизвестный код параметра", parameter_code=code or None, row=row_number))
            continue
        if code in values:
            errors.append(_issue("DUPLICATE_PARAMETER", "Параметр указан повторно", parameter_code=code, row=row_number))
            continue
        unit = (row.get("unit") or "").strip()
        if unit != parameter.unit:
            errors.append(_issue("UNIT_MISMATCH", f"Ожидается единица «{parameter.unit}»", parameter_code=code, row=row_number))
        value = _validate_value(parameter, row.get("value"), errors, row=row_number)
        if value is not None:
            values[code] = value
            locations[code] = {"row": row_number}
    _append_missing(parameters, values, errors)
    return values, errors, warnings, locations


def _check_xlsx_archive(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            infos = archive.infolist()
            if any('vbaproject' in item.filename.lower() or 'macrosheet' in item.filename.lower() for item in infos):
                raise IntakeError("XLSX_MACROS_FORBIDDEN")
            if len(infos) > 1000 or sum(item.file_size for item in infos) > MAX_XLSX_EXPANDED_BYTES:
                raise IntakeError("XLSX_ARCHIVE_LIMIT_EXCEEDED")
            if '[Content_Types].xml' in archive.namelist() and b'macroEnabled' in archive.read('[Content_Types].xml'):
                raise IntakeError('XLSX_MACROS_FORBIDDEN')
            if any(item.filename.startswith("/") or ".." in Path(item.filename).parts for item in infos):
                raise IntakeError("XLSX_ARCHIVE_INVALID")
    except (zipfile.BadZipFile, OSError, RuntimeError, NotImplementedError) as exc:
        raise IntakeError("XLSX_ARCHIVE_INVALID") from exc


def _parse_xlsx(
    payload: bytes, profile: ObjectProfile
) -> tuple[
    dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]
]:
    _check_xlsx_archive(payload)
    try:
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=True, keep_links=False
        )
    except (InvalidFileException, OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise IntakeError("XLSX_INVALID") from exc
    try:
        expected_sheet = profile.groups[0].parameters[0].source_sheet
        if expected_sheet not in workbook.sheetnames:
            raise IntakeError("XLSX_PROFILE_SHEET_MISSING")
        sheet = workbook[expected_sheet]
        expected_headers = ("Параметр", "Ед. изм.", "Базовое значение")
        actual_headers = tuple(sheet.cell(2, column).value for column in range(1, 4))
        if actual_headers != expected_headers:
            raise IntakeError("XLSX_HEADER_INVALID")
        values: dict[str, Any] = {}
        errors: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = []
        locations: dict[str, Any] = {}
        for parameter in profile.parameters():
            row = parameter.source_row
            label = sheet.cell(row, 1).value
            unit = sheet.cell(row, 2).value
            if label != parameter.label:
                errors.append(_issue("LABEL_MISMATCH", f"Ожидается параметр «{parameter.label}»", parameter_code=parameter.parameter_code, row=row))
                continue
            if unit != parameter.unit:
                errors.append(_issue("UNIT_MISMATCH", f"Ожидается единица «{parameter.unit}»", parameter_code=parameter.parameter_code, row=row))
            value = _validate_value(parameter, sheet.cell(row, 3).value, errors, row=row)
            if value is not None:
                values[parameter.parameter_code] = value
                locations[parameter.parameter_code] = {
                    "sheet": expected_sheet,
                    "row": row,
                    "cell": f"C{row}",
                }
        _append_missing(profile.parameters_by_code(), values, errors)
        return values, errors, warnings, locations
    finally:
        workbook.close()


def _append_missing(
    parameters: dict[str, ProfileParameter],
    values: dict[str, Any],
    errors: list[dict[str, Any]],
) -> None:
    already_reported = {
        issue.get("parameter_code")
        for issue in errors
        if issue["code"] in {"VALUE_MISSING", "LABEL_MISMATCH"}
    }
    for code in sorted(set(parameters) - set(values) - already_reported):
        errors.append(_issue("PARAMETER_MISSING", "Обязательная строка параметра отсутствует", parameter_code=code))


def inspect_project_file(
    filename: str,
    payload: bytes,
    profile_code: str,
) -> IntakeResult:
    safe_name = Path(filename or "").name
    if (
        not safe_name
        or safe_name != filename
        or "/" in filename
        or "\\" in filename
        or re.search(r"[\x00-\x1f]", safe_name)
    ):
        raise IntakeError("FILE_NAME_INVALID")
    if not payload:
        raise IntakeError("FILE_EMPTY")
    if len(payload) > MAX_FILE_BYTES:
        raise IntakeError("FILE_TOO_LARGE")
    try:
        profile = get_official_profile(profile_code)
    except ObjectProfileError as exc:
        raise IntakeError("PROFILE_UNSUPPORTED") from exc
    file_format = detect_format(safe_name, payload)
    from project_workbook import is_workbook, inspect_workbook
    if file_format == "XLSX":
        _check_xlsx_archive(payload)
    if is_workbook(payload, file_format):
        return inspect_workbook(safe_name, payload, profile.code, file_format)
    sha256 = hashlib.sha256(payload).hexdigest()
    if file_format == "XLSX":
        values, errors, warnings, locations = _parse_xlsx(payload, profile)
    else:
        values, errors, warnings, locations = _parse_csv(payload, profile)
    report = {
        "status": "VALID" if not errors else "INVALID",
        "parameter_count": len(profile.parameters()),
        "accepted_count": len(values),
        "error_count": len(errors),
        "warning_count": len(warnings),
        "errors": errors,
        "warnings": warnings,
    }
    file_source = {
        "name": safe_name,
        "sha256": sha256,
        "format": file_format,
    }
    parameter_provenance = {
        code: {
            "kind": "FILE",
            "parameter_code": code,
            "unit": profile.parameters_by_code()[code].unit,
            "source": {**file_source, **locations[code]},
        }
        for code in values
    }
    normalized_input = None
    provenance = None
    profile_version = ""
    if not errors:
        projection = build_profile_projection(
            profile,
            values,
            value_provenance_kind="FILE",
            file_source=file_source,
        )
        normalized_input = projection["normalized_input"]
        provenance = projection["provenance"]
        profile_version = projection["profile_version"]
    else:
        # The report still identifies the schema used for validation.
        projection = build_profile_projection(
            profile,
            {parameter.parameter_code: parameter.default_value for parameter in profile.parameters()},
            value_provenance_kind="PRESET",
        )
        profile_version = projection["profile_version"]
    return IntakeResult(
        file_format=file_format,
        profile_code=profile.code,
        profile_version=profile_version,
        original_name=safe_name,
        byte_size=len(payload),
        sha256=sha256,
        report=report,
        normalized_input=normalized_input,
        provenance=provenance,
        parameter_values=values,
        parameter_provenance=parameter_provenance,
    )


def build_csv_template(profile_code: str) -> tuple[str, bytes]:
    try:
        profile = get_official_profile(profile_code)
    except ObjectProfileError as exc:
        raise IntakeError("PROFILE_UNSUPPORTED") from exc
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(CSV_HEADERS)
    for parameter in profile.parameters():
        writer.writerow(
            (profile.code, parameter.parameter_code, parameter.default_value, parameter.unit)
        )
    return f"{profile.code}-project-input.csv", output.getvalue().encode("utf-8-sig")
