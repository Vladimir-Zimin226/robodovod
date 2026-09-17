from __future__ import annotations

import csv
import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from models import UserInput
from object_profiles import build_official_preset, get_official_profile
from project_file_intake import IntakeError, build_csv_template, inspect_project_file

SOURCE_XLSX = (
    Path(__file__).resolve().parents[1]
    / "Разобрать"
    / "Материалы от организаторов"
    / "Датасет"
    / "Датасеты_хакатон.xlsx"
)


def _profile_workbook(profile_code: str) -> bytes:
    profile = get_official_profile(profile_code)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = profile.groups[0].parameters[0].source_sheet
    sheet.cell(2, 1, "Параметр")
    sheet.cell(2, 2, "Ед. изм.")
    sheet.cell(2, 3, "Базовое значение")
    for parameter in profile.parameters():
        sheet.cell(parameter.source_row, 1, parameter.label)
        sheet.cell(parameter.source_row, 2, parameter.unit)
        sheet.cell(parameter.source_row, 3, parameter.default_value)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_xlsx_adapter_round_trip_without_restricted_fixture():
    result = inspect_project_file("warehouse.xlsx", _profile_workbook("warehouse"), "warehouse")

    assert result.valid
    assert result.report["accepted_count"] == 42
    assert len(result.parameter_values) == 42
    assert set(result.parameter_provenance) == set(result.parameter_values)
    assert {item["kind"] for item in result.parameter_provenance.values()} == {"FILE"}
    assert result.parameter_provenance["obschaya_ploschad_sklada"]["source"]["row"] == 4
    assert result.normalized_input == build_official_preset("warehouse")["normalized_input"]


@pytest.mark.parametrize(
    ("profile_code", "count"),
    [("warehouse", 42), ("airport", 39), ("medical_facility", 57)],
)
@pytest.mark.skipif(not SOURCE_XLSX.is_file(), reason="official XLSX is local-only")
def test_official_xlsx_equals_official_preset(profile_code: str, count: int):
    result = inspect_project_file(
        "Датасеты_хакатон.xlsx", SOURCE_XLSX.read_bytes(), profile_code
    )

    assert result.valid
    assert result.report == {
        "status": "VALID",
        "parameter_count": count,
        "accepted_count": count,
        "error_count": 0,
        "warning_count": 0,
        "errors": [],
        "warnings": [],
    }
    assert result.normalized_input == build_official_preset(profile_code)["normalized_input"]
    assert len(result.parameter_values) == count
    assert set(result.parameter_provenance) == set(result.parameter_values)
    assert {item["kind"] for item in result.provenance.values()} == {
        "FILE",
        "CALCULATED",
        "ASSUMED",
    }
    UserInput.model_validate(result.normalized_input)


@pytest.mark.parametrize(
    ("profile_code", "count"),
    [("warehouse", 42), ("airport", 39), ("medical_facility", 57)],
)
def test_generated_csv_round_trip(profile_code: str, count: int):
    filename, payload = build_csv_template(profile_code)
    result = inspect_project_file(filename, payload, profile_code)

    assert result.valid
    assert result.report["accepted_count"] == count
    assert len(result.parameter_values) == count
    assert result.normalized_input == build_official_preset(profile_code)["normalized_input"]


def test_csv_reports_unit_range_unknown_duplicate_and_missing_errors():
    profile = get_official_profile("warehouse")
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(("profile_code", "parameter_code", "value", "unit"))
    first = profile.parameters()[0]
    writer.writerow((profile.code, first.parameter_code, -1, "wrong"))
    writer.writerow((profile.code, first.parameter_code, first.default_value, first.unit))
    writer.writerow((profile.code, "unknown", 1, "шт."))

    result = inspect_project_file("bad.csv", output.getvalue().encode(), "warehouse")
    codes = {item["code"] for item in result.report["errors"]}

    assert not result.valid
    assert {"UNIT_MISMATCH", "BELOW_MIN", "DUPLICATE_PARAMETER", "UNKNOWN_PARAMETER", "PARAMETER_MISSING"} <= codes
    assert result.normalized_input is None


@pytest.mark.parametrize(
    ("name", "payload", "code"),
    [
        ("input.txt", b"x", "FILE_TYPE_UNSUPPORTED"),
        ("../input.csv", b"x", "FILE_NAME_INVALID"),
        ("input.xlsx", b"not-a-zip", "XLSX_SIGNATURE_INVALID"),
        ("input.csv", b"\xff", "CSV_ENCODING_INVALID"),
    ],
)
def test_malformed_files_fail_with_stable_code(name: str, payload: bytes, code: str):
    with pytest.raises(IntakeError, match=code):
        inspect_project_file(name, payload, "warehouse")
