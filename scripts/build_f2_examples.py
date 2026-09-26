"""Build reproducible synthetic acceptance inputs (never customer data)."""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from project_file_intake import inspect_project_file
from project_workbook import build_workbook


def build():
    out = ROOT / "docs/planning/assets/f2"
    out.mkdir(parents=True, exist_ok=True)
    fixtures = {}
    for profile in ["warehouse", "airport", "medical_facility"]:
        for demo in [False, True]:
            for extension in ["xlsx", "csv"]:
                name, payload = build_workbook(profile, demo, extension)
                (out / name).write_bytes(payload)
    for name, demand, batch in [
        ("interview-220-120", "220", "2"),
        ("typical-warehouse", "2000", "1"),
    ]:
        _, payload = build_workbook("warehouse", True)
        book = load_workbook(io.BytesIO(payload))
        for sheet in book:
            if sheet.title in {"Версия", "Инструкция"}:
                continue
            for row in sheet.iter_rows(min_row=2):
                code = row[1].value
                if sheet.title == "Процессы" and code == "demand":
                    row[3].value = demand
                if sheet.title == "Процессы" and code == "batch":
                    row[3].value = batch
                if sheet.title == "Зоны" and code == "constraints":
                    row[
                        3
                    ].value = "Синтетический склад: проход 3 м; пол ровный; груз 800 кг; обследование не выполнено"
                if sheet.title != "Паспорт" and row[5].value != "UNKNOWN":
                    row[5].value = (
                        "DATA" if name.startswith("interview") else "ASSUMPTION"
                    )
                    row[6].value = (
                        "Синтетический ответ для acceptance; не данные клиента"
                        if name.startswith("interview")
                        else "Типовой сценарий — допущение для acceptance"
                    )
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        filename = name + ".xlsx"
        (out / filename).write_bytes(stream.getvalue())
        result = inspect_project_file(filename, stream.getvalue(), "warehouse")
        assert result.valid, result.report
        fixtures[name] = result.normalized_input
    (out / "normalized-examples.json").write_text(
        json.dumps(fixtures, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return fixtures


if __name__ == "__main__":
    print("Built", list(build()))
