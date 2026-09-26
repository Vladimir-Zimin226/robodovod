"""F2 localhost browser journey. F2_DATABASE_URL must name a disposable stage11_* DB."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import subprocess
import tempfile
import time
import uuid
from pathlib import Path
from urllib.parse import urlparse

import psycopg
import requests
import websocket
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MOBILE = os.getenv("F2_MOBILE") == "1"
OUT = ROOT / (".tmp/f2/acceptance/mobile" if MOBILE else ".tmp/f2/acceptance")
OUT.mkdir(parents=True, exist_ok=True)
url = os.environ["F2_DATABASE_URL"].replace("postgresql+psycopg:", "postgresql:")
parsed = urlparse(url)
assert parsed.hostname in {"localhost", "127.0.0.1"} and parsed.path.startswith(
    "/stage11_"
)
API, WEB = "http://127.0.0.1:8000", "http://127.0.0.1:5173"


def snapshots():
    with psycopg.connect(url) as db:
        return {
            str(row[0]): row[1]
            for row in db.execute("SELECT id, row_to_json(r) FROM analysis_runs r")
        }


history = snapshots()
with psycopg.connect(url) as db:
    artifacts_before = {
        str(row[0]): row[1]
        for row in db.execute("SELECT id, row_to_json(r) FROM simulation_artifacts r")
    }
    old_email = db.execute(
        "SELECT email_normalized FROM users WHERE email_normalized LIKE 'stage31-%' ORDER BY created_at DESC LIMIT 1"
    ).fetchone()
old_owner = requests.Session()
old_export_url = None
if old_email:
    login = old_owner.post(
        API + "/api/auth/login",
        json={"email": old_email[0], "password": "Stage31-local-only-2026"},
        timeout=10,
    )
    assert login.ok
    projects = old_owner.get(API + "/api/projects", timeout=10).json()["items"]
    for old_project in projects:
        old_runs = old_owner.get(
            f"{API}/api/projects/{old_project['id']}/analysis-runs", timeout=10
        ).json()["items"]
        full = next(
            (row for row in old_runs if row["run_kind"] == "FULL_ANALYSIS"), None
        )
        if full:
            old_export_url = f"{API}/api/projects/{old_project['id']}/analysis-runs/{full['id']}/exports/report.pdf"
            export = old_owner.get(old_export_url, timeout=30)
            assert export.ok
            old_pdf = hashlib.sha256(export.content).hexdigest()
            break
session = requests.Session()
profile_path = Path(tempfile.mkdtemp(prefix="f2-chrome-", dir=ROOT / ".tmp"))
chrome = subprocess.Popen(
    [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        "--headless=new",
        "--disable-gpu",
        "--no-first-run",
        "--no-default-browser-check",
        "--remote-allow-origins=*",
        "--remote-debugging-port=0",
        f"--user-data-dir={profile_path}",
        "about:blank",
    ],
    stdout=subprocess.DEVNULL,
    stderr=subprocess.DEVNULL,
    creationflags=subprocess.CREATE_NO_WINDOW,
)
ws = None
seq = 0
browser_errors = []


def call(method, params=None):
    global seq
    seq += 1
    ws.send(json.dumps({"id": seq, "method": method, "params": params or {}}))
    while True:
        result = json.loads(ws.recv())
        if result.get("method") == "Runtime.exceptionThrown":
            browser_errors.append(result["params"])
        if result.get("id") == seq:
            assert "error" not in result, result
            return result.get("result", {})


def js(source):
    result = call(
        "Runtime.evaluate",
        {"expression": source, "returnByValue": True, "awaitPromise": True},
    )
    assert "exceptionDetails" not in result, result
    return result.get("result", {}).get("value")


def until(source, seconds=35):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        result = js(source)
        if result:
            return result
        time.sleep(0.2)
    raise AssertionError(f"{source}; body={js('document.body.innerText')}")


def click(label, selector="button"):
    assert js(
        "(()=>{const e=[...document.querySelectorAll(%s)].find(e=>e.textContent.trim().includes(%s));if(!e)return false;e.click();return true})()"
        % (json.dumps(selector), json.dumps(label))
    ), label


def screenshot(name):
    (OUT / name).write_bytes(
        base64.b64decode(
            call("Page.captureScreenshot", {"captureBeyondViewport": False})["data"]
        )
    )


def get_project(pid):
    result = session.get(f"{API}/api/projects/{pid}", timeout=10)
    assert result.ok, result.text
    return result.json()


def post(route, body):
    result = session.post(API + route, json=body, headers=headers, timeout=45)
    assert result.ok, (route, result.status_code, result.text)
    return result.json()


try:
    until_file = profile_path / "DevToolsActivePort"
    for _ in range(80):
        if until_file.exists():
            break
        time.sleep(0.2)
    port = until_file.read_text().splitlines()[0]
    pages = requests.get(f"http://127.0.0.1:{port}/json/list", timeout=5).json()
    ws = websocket.create_connection(
        next(p["webSocketDebuggerUrl"] for p in pages if p["type"] == "page"),
        timeout=10,
    )
    call("Runtime.enable")
    call("Page.enable")
    call(
        "Emulation.setDeviceMetricsOverride",
        {
            "width": 390 if MOBILE else 1366,
            "height": 844 if MOBILE else 768,
            "deviceScaleFactor": 1,
            "mobile": MOBILE,
        },
    )
    downloads = OUT / "downloads"
    downloads.mkdir(exist_ok=True)
    call(
        "Page.setDownloadBehavior",
        {"behavior": "allow", "downloadPath": str(downloads)},
    )
    call("Page.navigate", {"url": WEB + "/#templates"})
    until('document.body.innerText.includes("Пустой XLSX")')
    assert not js('Boolean(document.querySelector("input[type=file]"))')
    click("Пустой XLSX", "a")
    for _ in range(60):
        paths = list(downloads.glob("*blank.xlsx"))
        if paths:
            break
        time.sleep(0.2)
    book = load_workbook(paths[0], read_only=True)
    assert book["Версия"]["B1"].value == "project-workbook-v1"
    book.close()
    for profile in ["warehouse", "airport", "medical_facility"]:
        for variant in ["blank", "demo"]:
            result = requests.get(
                f"{API}/api/project-workbooks/{profile}/{variant}.xlsx", timeout=10
            )
            assert result.ok
            book = load_workbook(io.BytesIO(result.content))
            assert len(book.sheetnames) == 8
            book.close()
    screenshot("guest-desktop.png")
    call(
        "Emulation.setDeviceMetricsOverride",
        {"width": 390, "height": 844, "deviceScaleFactor": 1, "mobile": True},
    )
    screenshot("guest-mobile.png")
    assert js("document.documentElement.scrollWidth <= innerWidth"), "mobile overflow"
    call(
        "Emulation.setDeviceMetricsOverride",
        {
            "width": 390 if MOBILE else 1366,
            "height": 844 if MOBILE else 768,
            "deviceScaleFactor": 1,
            "mobile": MOBILE,
        },
    )
    auth = session.post(
        API + "/api/auth/register",
        json={
            "email": f"f2-{uuid.uuid4().hex[:12]}@example.com",
            "password": "F2-local-only-2026",
            "name": "F2 acceptance",
        },
        timeout=10,
    )
    assert auth.status_code == 201, auth.text
    headers = {"X-CSRF-Token": auth.json()["csrf_token"]}
    for cookie in session.cookies:
        call(
            "Network.setCookie",
            {"name": cookie.name, "value": cookie.value, "url": WEB},
        )
    results = []
    for example in ["interview-220-120", "typical-warehouse"]:
        project = post("/api/projects", {"name": "F2 · " + example})
        pid = project["id"]
        js(
            f"localStorage.setItem({json.dumps('robodovod.active-project.' + auth.json()['user']['id'])}, {json.dumps(pid)})"
        )
        call("Page.navigate", {"url": WEB + f"/?f2={pid}#templates"})
        until(
            "Boolean(document.querySelector('main[aria-label=\"Шаблоны и загрузка данных\"] select'))"
        )
        until(
            f'[...document.querySelectorAll("select option")].some(e=>e.value==={json.dumps(pid)})'
        )
        js(
            '(()=>{const e=[...document.querySelectorAll("main select")].find(e=>[...e.options].some(o=>o.value===%s));e.value=%s;e.dispatchEvent(new Event("change",{bubbles:true}))})()'
            % (json.dumps(pid), json.dumps(pid))
        )
        until('Boolean(document.querySelector("input[type=file]"))')
        node = call("DOM.getDocument")["root"]["nodeId"]
        file_node = call(
            "DOM.querySelector", {"nodeId": node, "selector": "input[type=file]"}
        )["nodeId"]
        call(
            "DOM.setFileInputFiles",
            {
                "nodeId": file_node,
                "files": [str(ROOT / f"docs/planning/assets/f2/{example}.xlsx")],
            },
        )
        until(f"document.body.innerText.includes({json.dumps(example + '.xlsx')})")
        before = get_project(pid)
        click("Проверить", ".project-file-intake button")
        until('document.body.innerText.includes("Файл прошёл проверку")')
        assert get_project(pid) == before
        assert js('document.body.innerText.includes("Источники и предложения")')
        screenshot(example + "-preview.png")
        click("Применить к базовому сценарию")
        until('document.body.innerText.includes("Проверить и подтвердить в форме")')
        project = get_project(pid)
        assert all(
            not item["confirmed_by_user"]
            for row in project["profile"]["file_intake_v2"]["records"][
                "Процессы"
            ].values()
            for item in row.values()
        )
        click("Проверить и подтвердить в форме")
        until('document.body.innerText.includes("Подтвердить входы книги для расчёта")')
        click("Подтвердить входы книги для расчёта")
        click("Проверить ввод")
        until(
            '[...document.querySelectorAll("aside label")].find(e=>e.textContent.includes("Модель из активного capacity-каталога"))?.querySelector("select")?.options.length > 1'
        )
        until(
            'document.body.innerText.includes("Модель из активного capacity-каталога")'
        )
        js(
            '(()=>{const e=[...document.querySelectorAll("aside label")].find(e=>e.textContent.includes("Модель из активного capacity-каталога")).querySelector("select");e.value=e.options[1].value;e.dispatchEvent(new Event("change",{bubbles:true}))})()'
        )
        js(
            '[...document.querySelectorAll("aside label")].find(e=>e.textContent.includes("Подтверждаю, что данные типового")).querySelector("input").click()'
        )
        click("Рассчитать и сохранить")
        until('Boolean(document.querySelector(".capacity-results-v2"))', 45)
        until('Boolean(document.querySelector("form.economics-inputs-v2"))')
        assert js(
            'document.querySelector("#economics-manual_units_per_shift")?.value === "100"'
        )
        assert js(
            '[...document.querySelectorAll("form.economics-inputs-v2 input[type=checkbox]")].every(e=>!e.checked)'
        )
        run_list = session.get(
            f"{API}/api/projects/{pid}/analysis-runs", timeout=10
        ).json()["items"]
        capacity = next(
            item for item in run_list if item["run_kind"] == "CAPACITY_ANALYSIS"
        )
        saved = session.get(
            f"{API}/api/projects/{pid}/analysis-runs/{capacity['id']}", timeout=10
        ).json()
        # Manual draft is reconstructed independently through the public form helpers.
        manual = js(
            """(async()=>{const m=await import('/src/processRoleIntakeV2.js');let d=m.createDraft('retail');
          const input=%s;const p=input.records['Процессы'].operation; const role=input.records['Роли'].staff;
          d=m.updateProcess(d,'warehouse_receiving_shipping',{active:true,demand:p.demand.value,shifts:p.shifts.value,hours:p.hours.value,days:p.days.value,distance:p.distance.value,batch:p.batch.value});
          d=m.setRoleActive(d,'warehouse_receiving_shipping','forklift_driver',true);d=m.updateRole(d,'draft.warehouse.forklift_driver',{headcount:role.headcount.value,salary:role.salary.value});
          const request=m.serializeDraft(d);const response=await fetch('/api/v2/calculation-intake/normalize',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(request)});return {request,response:await response.json()}})()"""
            % json.dumps(project["profile"]["file_intake_v2"])
        )
        assert manual["response"]["valid"], manual
        imported = saved["input_snapshot"]["process"]
        hand = next(
            p for p in manual["response"]["normalized_processes"] if p["active"]
        )
        for key in ["demand", "route_distance", "explicit_batch"]:
            assert imported[key]["normalized_value"] == hand[key]["normalized_value"], (
                key,
                imported,
                hand,
            )
        manual_request = json.loads(json.dumps(saved["input_snapshot"]))
        manual_request.update(
            process=hand,
            input_revision=hand["input_revision"],
            role_pool=manual["response"]["role_pool"],
        )
        manual_request["process"]["exchange"] = imported["exchange"]
        manual_request["zone_context"]["zone_id"] = "zone.draft.warehouse.main"
        manual_result = post("/api/v2/capacity-analyses", manual_request)
        (OUT / f"{example}-comparison.json").write_text(
            json.dumps(
                {
                    "imported": saved,
                    "manual": manual_result,
                    "manual_request": manual_request,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        assert (
            manual_result["capacity"]["value"]
            == saved["result_snapshot"]["capacity"]["value"]
        )
        assert (
            manual_result["capacity"]["status"]
            == saved["result_snapshot"]["capacity"]["status"]
        )
        click("Подтвердить допущения", "form.economics-inputs-v2 button")
        for selector in [
            "grossConfirm",
            "currencyConfirm",
            "initialBatteryConfirm",
            "batteryServiceConfirm",
            "raasScopeConfirm",
        ]:
            assert js(
                f'(()=>{{const e=document.querySelector("#economics-condition-{selector}");if(!e)return false;e.click();return true}})()'
            ), selector
        click("Сохранить частичный результат", "form.economics-inputs-v2 button")
        until(
            'Boolean(document.querySelector(\'[aria-label="Коммерческие сценарии"], [aria-label="Частичный результат экономики"]\'))',
            45,
        )
        screenshot(example + "-result.png")
        results.append(
            {
                "example": example,
                "project_id": pid,
                "capacity_run_id": saved["id"],
                "capacity": saved["result_snapshot"]["capacity"]["value"],
                "manual_equal": True,
            }
        )
        # Brain gets the same proposal and requires its own explicit confirmation.
        call("Page.navigate", {"url": WEB + f"/?f2={pid}#model"})
        until('document.body.innerText.includes("Паспорт и источники книги")')
        brain = session.get(f"{API}/api/brain/projects/{pid}", timeout=10).json()
        assert (
            not brain["profile"]["preflight_confirmed"]
            and brain["profile"]["imported_workbook"]
            == project["profile"]["file_intake_v2"]
        )
        click("Подтвердить профиль", '[aria-label="Профиль расчёта"] button')
        until(
            f'document.body.innerText.includes("версия профиля: {brain["profile"]["profile_version"] + 1}")'
        )
        until(
            "[...document.querySelectorAll('[aria-label=\"Расчёт и результат\"] select option')].some(e=>e.value==="
            + json.dumps(saved["input_snapshot"]["position_id"])
            + ")"
        )
        js(
            '(()=>{const e=document.querySelector(\'[aria-label="Расчёт и результат"] select\');e.value=%s;e.dispatchEvent(new Event("change",{bubbles:true}))})()'
            % json.dumps(saved["input_snapshot"]["position_id"])
        )
        click("Рассчитать и сохранить", '[aria-label="Расчёт и результат"] button')
        until('Boolean(document.querySelector(".capacity-results-v2"))', 45)
        linked = session.get(f"{API}/api/brain/projects/{pid}", timeout=10).json()[
            "profile"
        ]["runs"]
        brain_run = next(
            item for item in linked if item["run_kind"] == "CAPACITY_ANALYSIS"
        )
        result = session.get(
            f"{API}/api/projects/{pid}/analysis-runs/{brain_run['run_id']}", timeout=10
        ).json()
        assert (
            result["result_snapshot"]["capacity"]["value"]
            == saved["result_snapshot"]["capacity"]["value"]
        )
        results[-1]["brain_equal"] = True
    after = snapshots()
    assert all(after[rid] == snapshot for rid, snapshot in history.items())
    with psycopg.connect(url) as db:
        artifacts_after = {
            str(row[0]): row[1]
            for row in db.execute(
                "SELECT id, row_to_json(r) FROM simulation_artifacts r"
            )
        }
    assert all(artifacts_after[rid] == row for rid, row in artifacts_before.items())
    if old_export_url:
        assert (
            hashlib.sha256(
                old_owner.get(old_export_url, timeout=30).content
            ).hexdigest()
            == old_pdf
        )
    assert not browser_errors, browser_errors
    (OUT / "summary.json").write_text(
        json.dumps(
            {
                "results": results,
                "historical_runs_unchanged": len(history),
                "historical_simulation_artifacts_unchanged": len(artifacts_before),
                "historical_pdf_sha256": old_pdf if old_export_url else None,
                "browser_errors": browser_errors,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print("F2 browser passed", results, "historical runs unchanged:", len(history))
finally:
    if ws:
        ws.close()
    chrome.terminate()
    chrome.wait(timeout=10)
