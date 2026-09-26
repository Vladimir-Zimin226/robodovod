"""Generate reviewable DOCX/PPTX and current OpenAPI without credentials."""

from __future__ import annotations

import importlib.metadata
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "delivery" / "f8"
sys.path.insert(0, str(ROOT / "backend"))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    code = False
    for line in (ROOT / "docs" / "FINAL_DELIVERY_2026-09-27.md").read_text(encoding="utf-8").splitlines():
        if line.startswith("```"):
            code = not code
        elif line.startswith("# "):
            document.add_heading(line[2:], 0)
        elif line.startswith("## "):
            document.add_heading(line[3:], 1)
        elif line:
            paragraph = document.add_paragraph(line)
            if code:
                for run in paragraph.runs:
                    run.font.name = "Consolas"
    document.save(OUT / "Robodovod-user-admin-technical-guide.docx")
    slides = [
        ("РОБОДОВОД", "Предварительное ТЭО роботизации\nСклад: полный путь от входов до отчёта\n27.09.2026 · локальный кандидат выпуска"),
        ("Задача", "Сопоставить потребный парк и затраты\nПоказать источники и неизвестные данные\nСохранить воспроизводимые версии результата"),
        ("Ввод", "Ручная форма, Brain или XLSX/CSV\n220 паллет/сутки · 120 м в одну сторону\nЯвные график, пакет и обмен; отдельное подтверждение"),
        ("Подбор", "Каталог по объекту и процессу\nINCLUDED / EXCLUDED / REQUIRES_CHECK\nНепроверенные ТТХ не превращаются в PASS"),
        ("Экономика", "Baseline / покупка / RaaS на общем горизонте\nCAPEX, OPEX, TCO, NPV, окупаемость\nROI с явной базой; sensitivity ±10%"),
        ("Физический сценарий", "ScenarioSpec и сохранённый C23\nПоток, очередь и график\n2D/3D с общей привязкой; геометрия условная"),
        ("Доказательства", "Immutable input/result/trace и checksums\nPDF, XLSX, CSV, SVG, ZIP\nReopen и новая версия без переписывания истории"),
        ("Администратор", "Новый черновик каталога\nТТХ, цены, источники, нормы\nValidate → publish → отдельная activation"),
        ("Три объекта", "Склад — полный расчётный путь\nАэропорт и клиника — паспорт и релевантный каталог\nНеподдержанные физические процессы обозначены"),
        ("Надёжность и границы", "Демо и числовое ядро без YC\nBrain: сохранение до вызова и fallback\nНужны обследование, коммерческое подтверждение и HTTPS-приёмка"),
        ("Показ — 7 минут", "Демо → проект 220/120 → каталог → C11\nЭкономика → C23 → 2D/3D → экспорт\nИстория → ADMIN каталог → ограничения"),
    ]
    presentation = Presentation()
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    presentation.save(OUT / "Robodovod-final-presentation.pptx")
    import main as backend
    (OUT / "openapi.json").write_text(json.dumps(backend.app.openapi(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    packages = []
    for name in ("fastapi", "pydantic", "SQLAlchemy", "alembic", "requests", "psycopg", "argon2-cffi", "pypdf", "Pillow", "openpyxl"):
        metadata = importlib.metadata.metadata(name)
        packages.append({"name": name, "installed_version": metadata["Version"],
                         "license": metadata.get("License-Expression") or metadata.get("License") or "see distribution license files"})
    (OUT / "python-license-inventory.json").write_text(json.dumps(packages, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lock = json.loads((ROOT / "frontend" / "package-lock.json").read_text(encoding="utf-8"))
    javascript = [{"name": name.removeprefix("node_modules/"), "version": details.get("version"),
                   "license": details.get("license", "check package distribution")}
                  for name, details in lock["packages"].items() if name]
    (OUT / "javascript-license-inventory.json").write_text(json.dumps(javascript, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
