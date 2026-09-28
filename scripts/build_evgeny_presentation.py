"""Build the editable E9 deck from the organizer's PowerPoint template.

The template and organizer screenshots are read only. Browser captures are
copied to the delivery folder before running this script.
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ORGANIZER = ROOT / "Разобрать/Материалы от организаторов/Презентация"
OUT = ROOT / "docs/delivery/evgeny-feedback-2026-09-28"
ASSETS = OUT / "presentation-assets"
DEST = OUT / "Robodovod-ZMNCRAFT-2026.pptx"
P = RGBColor(49, 15, 83)
PINK = RGBColor(255, 0, 83)
LILAC = RGBColor(138, 131, 209)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(28, 29, 34)
PALE = RGBColor(255, 214, 228)


def box(slide, x, y, w, h, color=WHITE, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def label(slide, value, x, y, w, h, size=21, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(7)
        if align:
            paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = "Montserrat"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def panel(slide, title, subtitle, number):
    # The original slide, theme, background and organizer logos stay in place.
    box(slide, .34, 1.06, 12.65, 5.98, WHITE, radius=True)
    label(slide, title, .72, 1.30, 11.8, .55, 29, P, True)
    box(slide, .72, 1.96, 1.25, .06, PINK)
    label(slide, subtitle, .73, 2.08, 11.8, .40, 13, LILAC)
    label(slide, f"ZMNCRAFT · РОБОДОВОД  •  {number:02d}", .72, 6.63, 11.8, .23, 10, P)


def card(slide, heading, body, x, y, w, h, heading_size=19, body_size=15):
    box(slide, x, y, w, h, PALE, radius=True)
    label(slide, heading, x + .20, y + .15, w - .40, .42, heading_size, P, True)
    label(slide, body, x + .20, y + .66, w - .40, h - .78, body_size, INK)


def picture(slide, path, x, y, w, h):
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    ww, hh = iw * scale, ih * scale
    left, top = x + (w - ww) / 2, y + (h - hh) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(ww), Inches(hh))


def main():
    prs = Presentation(ORGANIZER / "ЛЦТ2026 Шаблон презентации.pptx")
    keep = list(range(6, 20))  # organizer slides 7–20, 14 finished slides
    ids = prs.slides._sldIdLst
    for index in reversed(range(len(ids))):
        if index not in keep:
            ids.remove(ids[index])
    slides = list(prs.slides)
    for template_slide in slides:
        for shape in template_slide.shapes:
            if shape.has_text_frame:
                shape.text = ""

    # 1. Organizer's title layout and marks.
    slide = slides[0]
    box(slide, .55, 2.06, 8.35, 3.58, P, radius=True)
    label(slide, "РОБОДОВОД", .84, 2.38, 7.5, .83, 43, WHITE, True)
    label(slide, "ZMNCRAFT · ЛЦТ 2026", .87, 3.23, 7.4, .42, 20, PALE, True)
    label(slide, "Платформа подбора роботизированных решений с расчётом экономического эффекта и визуализацией работы роботов на объекте", .87, 3.80, 7.42, 1.40, 18, WHITE)
    label(slide, "Задача Лидеров цифровой трансформации · 28 сентября 2026", .86, 6.66, 10.0, .28, 12, WHITE)

    slide = slides[1]
    panel(slide, "Проблема и решение", "Предварительная инвестиционная оценка до обследования объекта", 2)
    card(slide, "Проблема", "Робота легко выбрать по каталогу, но сложнее проверить нагрузку, смены, маршрут, исходный труд и расходы проекта.", .72, 2.72, 5.85, 3.20)
    card(slide, "Решение", "РОБОДОВОД связывает проверку условий, подбор парка, труд до/после, покупку или RaaS и 2D/3D показ одного сохранённого сценария.", 6.76, 2.72, 5.84, 3.20)

    slide = slides[2]
    panel(slide, "Команда ZMNCRAFT", "Два участника · контакты для связи", 3)
    for x, photo, name, role, details, tg in [
        (.75, "Зимин.png", "Зимин Владимир Андреевич", "Капитан · AI-product / full-stack", "Архитектура и продуктовая логика\nМагистрант ИТМО · Томари", "@vovzmncraft"),
        (6.80, "Юрченко.png", "Юрченко Евгений Александрович", "AI-инженер · vibe-кодер", "Расчётная логика, frontend и backend\nЮрист · университет Хабаровска", "@pawuk_ptr"),
    ]:
        box(slide, x, 2.65, 5.78, 3.57, PALE, radius=True)
        picture(slide, ORGANIZER / photo, x + .24, 2.92, 1.45, 1.45)
        label(slide, name, x + 1.86, 2.85, 3.56, .72, 16, P, True)
        label(slide, role, x + 1.86, 3.55, 3.55, .67, 15, INK)
        label(slide, details, x + .24, 4.58, 5.25, .86, 15, INK)
        label(slide, f"Telegram: {tg}", x + .24, 5.57, 5.23, .36, 16, P, True)

    slide = slides[3]
    panel(slide, "Как команда решала задачу", "История знакомства участников не документирована; показан подтверждённый рабочий вклад", 4)
    card(slide, "Мотивация продукта", "Дать заказчику проверяемый первый ответ: сколько роботов нужно, какой труд остаётся и что меняется в денежном результате.", .72, 2.68, 5.85, 2.90)
    card(slide, "Сложность и подход", "Сохранили старые результаты без пересчёта. Отделили рабочую долю, зарядку и новый штат; неизвестные цены оставили частичным результатом.", 6.77, 2.68, 5.82, 2.90)

    slide = slides[4]
    panel(slide, "Техническая и рыночная суть", "Инструмент для предварительного разговора заказчика с интегратором", 5)
    card(slide, "Техническая", "React/Vite + FastAPI/PostgreSQL. C05 проверяет ограничения, C11 парк, C14 труд, C18 проектные потоки; версии и checksum связывают результаты и PDF.", .72, 2.70, 5.85, 3.08)
    card(slide, "Практическая", "Склад, аэропорт и клиника — типовые входы для демонстрации. Сравнение покупки/RaaS помогает определить данные для пилота. Продажи и внедрения не заявлены.", 6.77, 2.70, 5.83, 3.08)

    slide = slides[5]
    panel(slide, "Путь пользователя", "Собственный скриншот обновлённого локального сервиса · типовой ввод", 6)
    picture(slide, ASSETS / "warehouse-intake.png", .74, 2.58, 7.42, 3.58)
    card(slide, "От объекта к результату", "1. Типовой объект или свои данные\n2. Проверка источников и C05\n3. Парк и штат\n4. Экономика и сохранение\n5. Симуляция, what-if и PDF", 8.25, 2.58, 4.31, 3.58, 18, 15)

    slide = slides[6]
    panel(slide, "Парк привязан к операции", "Контрольный склад: паллетная перевозка, не автономный отбор коробок", 7)
    card(slide, "2 000 паллет/сутки", "Плечо 120 м · тестовый график 2 × 10 ч\nРасчётная потребность: 20 роботов\nИсточник: изолированный контрольный профиль E1", .73, 2.70, 5.86, 3.05)
    card(slide, "220 паллет/сутки", "То же плечо и график\nРасчётная потребность: 3 робота\nПарк меняется с объёмом, а не от названия склада", 6.74, 2.70, 5.85, 3.05)

    slide = slides[7]
    panel(slide, "Труд до и после", "Высвобождение задачи не означает увольнение людей", 8)
    card(slide, "Исходный штат", "Вводятся роли, численность и gross зарплаты. Доля роботизируемой работы ограничивает возможное высвобождение; остаточные операции остаются ручными.", .73, 2.69, 5.86, 3.20)
    card(slide, "Новые функции", "Диспетчеры и техники считаются отдельно по парку, сменам и ротации. Перевод, найм и сервис поставщика различаются в C14 и расходах.", 6.76, 2.69, 5.82, 3.20)

    slide = slides[8]
    panel(slide, "Покупка и RaaS", "Единая проектная база C18 · финансовый вывод для каждой вкладки", 9)
    card(slide, "Покупка", "Цена оборудования × парк, внедрение суммой или процентом, эксплуатация, новый штат и потоки по годам.", .72, 2.71, 5.85, 3.04)
    card(slide, "RaaS", "Месячная ставка на робота суммой или процентом от gross цены. Сервисный персонал учитывается по выбранному договорному режиму.", 6.76, 2.71, 5.82, 3.04)
    label(slide, "NPV и сроки окупаемости зависят от сохранённого сценария; неподтверждённая закупка остаётся ограничением.", .78, 6.06, 11.73, .36, 15, P)

    slide = slides[9]
    panel(slide, "Глубина и «Что, если»", "Денежный preview использует серверный движок и не изменяет исходный run", 10)
    card(slide, "Базовый → Углублённый → Полный", "Уровни раскрывают труд, коммерческие варианты и проектные потоки. PDF отмечает глубину на каждой странице.", .72, 2.68, 5.85, 2.88)
    card(slide, "Новая версия расчёта", "Изменение цены, внедрения или RaaS сравнивается с исходником. Сохранение создаёт отдельный immutable run; спрос и окно требуют нового C11.", 6.77, 2.68, 5.82, 2.88)

    slide = slides[10]
    panel(slide, "Сценарий в 2D и 3D", "Скриншот локальной проверки сохранённого примера · авторский профиль", 11)
    picture(slide, ASSETS / "guest-simulation.png", .75, 2.56, 8.45, 3.62)
    card(slide, "Один источник", "Плеер показывает текущий run. Переключение вида не подменяет физический сценарий другим проектом.", 9.34, 2.66, 3.21, 3.20, 17, 14)

    slide = slides[11]
    panel(slide, "Архитектура и доказуемость", "Из сохранённого ввода в PDF без тихого пересчёта", 12)
    for x, heading, body in [(.73, "Ввод", "Процесс, объект, роли, график и источники"), (4.75, "Расчёт", "C05 → C11 → C14 → C18 → C23"), (8.77, "Экспорт", "Immutable run, checksum, PDF и ZIP")]:
        card(slide, heading, body, x, 2.72, 3.80, 2.82, 19, 16)
    label(slide, "Исторические расчёты не переписываются новой методикой. Backup использован только для чтения.", .77, 5.78, 11.70, .45, 15, P)

    slide = slides[12]
    panel(slide, "Пилот и развитие", "Следующий шаг — подтвердить физику и цены на объекте", 13)
    card(slide, "Пилот", "Замерить маршруты, спрос, окна, покрытия, зарядку и остаточные операции; получить паспорт робота и предложение поставщика.", .73, 2.69, 5.84, 3.00)
    card(slide, "Развитие", "Физический what-if от нового C11; сохранённый профиль комплектации PICK/ORDER_LINE; подробные факторы C19 и проверенная интеграция с БД.", 6.75, 2.69, 5.83, 3.00)

    slide = slides[13]
    panel(slide, "Материалы и контакты", "Документация и код в репозитории; доступность внешнему проверяющему зависит от публикации", 14)
    label(slide, "Прототип: robodovod.ru\nРепозиторий: github.com/Vladimir-Zimin226/robodovod\nДокументация: docs/delivery/evgeny-feedback-2026-09-28/\nКоманда: zmncraft.ru\nTelegram: @vovzmncraft · @pawuk_ptr", .80, 2.67, 11.7, 2.66, 19, P)
    label(slide, "Предварительная оценка. Авторские допущения типовых объектов не являются коммерческим предложением или подтверждением внедрения.", .80, 5.75, 11.68, .63, 15, INK)

    prs.save(DEST)
    print(f"{DEST}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
