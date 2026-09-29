"""Export the current editable ZMNCRAFT deck to PDF using local PowerPoint.

Run locally from the repository root: python scripts/export_story_presentation_pdf.py
PowerPoint is required; this script never changes the PPTX.
"""

from pathlib import Path
from hashlib import sha256

import win32com.client
from pypdf import PdfReader
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / 'docs/presentation/Robodovod-ZMNCRAFT-2026.pptx'
PDF = PPTX.with_suffix('.pdf')
PENDING_PDF = PDF.with_name(PDF.stem + '.new.pdf')
SOURCES = PPTX.parent / 'SOURCES.md'
PENDING = '**PDF сейчас от прежней 20-слайдовой версии и не соответствует этой PPTX.** Его требуется пересоздать установленным PowerPoint командой `python scripts/export_story_presentation_pdf.py` и проверить, что в нём 18 страниц. Экспорт текущей версии был остановлен автоматической проверкой доступа к PowerPoint из-за исчерпанного лимита инструмента; это не отказ по безопасности.'


def main():
    source = Presentation(PPTX)
    if len(source.slides) != 18:
        raise RuntimeError(f'Expected 18 PPTX slides, found {len(source.slides)}')
    if not source.slides[0].shapes[-1]._element.xpath('.//*[local-name()="svgBlip"]'):
        raise RuntimeError('The ZMNCRAFT SVG logo is missing from slide 1')
    for index, slide in enumerate(source.slides, 1):
        if not any(shape.name == 'Unified page number' and shape.text == f'{index:02d} / 18'
                   for shape in slide.shapes if shape.has_text_frame):
            raise RuntimeError(f'Slide {index} has no matching page number')
    app = win32com.client.DispatchEx('PowerPoint.Application')
    deck = None
    try:
        deck = app.Presentations.Open(str(PPTX), True, False, False)
        if deck.Slides.Count != 18:
            raise RuntimeError(f'Expected 18 slides, found {deck.Slides.Count}')
        deck.SaveAs(str(PENDING_PDF), 32)
        if not PENDING_PDF.exists() or PENDING_PDF.stat().st_size < 100_000:
            raise RuntimeError('PDF export did not produce a complete file')
        pages = len(PdfReader(PENDING_PDF).pages)
        if pages != deck.Slides.Count:
            raise RuntimeError(f'PDF has {pages} pages, PPTX has {deck.Slides.Count} slides')
        PENDING_PDF.replace(PDF)
        digest = sha256(PDF.read_bytes()).hexdigest().upper()
        source_digest = sha256(PPTX.read_bytes()).hexdigest().upper()
        if SOURCES.exists():
            text = SOURCES.read_text(encoding='utf-8')
            text = text.replace(PENDING, '**PDF обновлён из текущей PPTX через локальный PowerPoint и содержит 18 страниц.**')
            if '| `Robodovod-ZMNCRAFT-2026.pdf` |' not in text:
                text += f'| `Robodovod-ZMNCRAFT-2026.pdf` | `{digest}` |\n'
            else:
                import re
                text = re.sub(r'\| `Robodovod-ZMNCRAFT-2026\.pdf` \| `[^`]+` \|',
                              f'| `Robodovod-ZMNCRAFT-2026.pdf` | `{digest}` |', text)
            import re
            text = re.sub(r'\| `Robodovod-ZMNCRAFT-2026\.pptx` \| `[^`]+` \|',
                          f'| `Robodovod-ZMNCRAFT-2026.pptx` | `{source_digest}` |', text)
            SOURCES.write_text(text, encoding='utf-8')
        print(f'{PDF}: {pages} pages, {PDF.stat().st_size} bytes, SHA-256 {digest}')
    finally:
        if deck is not None:
            deck.Close()
        app.Quit()


if __name__ == '__main__':
    main()
